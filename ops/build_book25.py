"""
build_book25.py — the 25-slot managed-book workbook (HANDOVER 3.102).

A fresh workbook (no trade history) for the design in ops/book25_design.md:
12 Core, 5 Bench, 8 Shadow slots; breadth gate with held cash lent to the Bench
(option E); 25% exposure cap and per-name liquidity cap; Off holds cash; a
Shadow admission panel, a daily monitoring Dashboard, a Parameters panel with
review cadence, a Change Log, and step-by-step procedures in the Guide.

Slot-based: every slot has the same fixed rows, columns and Model sheet
(M01..M25), so changing a slot's ticker or tier never moves the layout the
scripts depend on.

The Model sheets reproduce the live workbook's per-name engine (Model TSM)
formula for formula; only their data source (Query, by slot) and their
parameter cells (Slots / Parameters) change.

Usage:
    python3 ops/build_book25.py <live_workbook.xlsx> <seed.json> <out.xlsx>
"""
import datetime as dt
import json
import re
import sys

import openpyxl
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

NS = 25                     # slots
N = 1260                    # data rows of capacity (~5 years from Apr 2024)
Q0, Q1 = 2, N + 1           # Query data rows
M0, M1 = 8, N + 7           # Model data rows (Model row r <-> Query row r-6)
QF = ('O', 'H', 'L', 'C', 'DV', 'FH')   # Query fields per slot
SL0 = 7                     # Slots: first slot row
LV0 = 22                    # Engine: first levels row
SV0 = 51                    # Engine: first sleeve row
OR0 = 9                     # Orders: first order row
LG0, LG1 = 7, 2006          # Log: data rows
NAV0, NAV1 = 34, 1533       # Performance: NAV log rows

# ----------------------------------------------------------------- styling
FONT = 'Calibri'
NAVY, BLUE, LIGHT, PALE, GREY = '1F3864', '2F5597', 'D9E1F2', 'EEF2F9', 'F2F2F2'
INPUT_FILL, INPUT_FONT = 'FFF2CC', '0000FF'
SCRIPT_FILL = 'E2EFDA'
RED_F, RED_T = 'FFC7CE', '9C0006'
AMB_F, AMB_T = 'FFEB9C', '9C5700'
GRN_F, GRN_T = 'C6EFCE', '006100'
TIER = {'Core': ('1F3864', 'FFFFFF', False), 'Bench': ('BDD7EE', '1F4E79', False),
        'Shadow': ('EDEDED', '595959', True), 'Off': ('FCE4D6', '843C0C', False)}
THIN = Side(style='thin', color='BFBFBF')
HAIR = Side(style='hair', color='D9D9D9')


def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)


def font(size=10, bold=False, color='000000', italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def put(ws, ref, value, *, f=None, fl=None, fmt=None, al=None, bd=None):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fl:
        c.fill = fill(fl)
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if bd:
        c.border = bd
    return c


def title(ws, text, sub):
    ws.sheet_view.showGridLines = False
    put(ws, 'B1', text, f=font(16, True, NAVY))
    put(ws, 'B2', sub, f=font(10, False, '595959', True))
    ws.row_dimensions[1].height = 26
    ws.column_dimensions['A'].width = 2


def bar(ws, row, c1, c2, text):
    for c in range(c1, c2 + 1):
        ws.cell(row, c).fill = fill(NAVY)
    put(ws, f'{CL(c1)}{row}', text, f=font(10, True, 'FFFFFF'))
    ws.row_dimensions[row].height = 18


def header(ws, row, c1, labels, height=30):
    for i, lab in enumerate(labels):
        put(ws, f'{CL(c1 + i)}{row}', lab, f=font(9, True, NAVY), fl=LIGHT,
            al=Alignment(wrap_text=True, vertical='center', horizontal='center'),
            bd=Border(bottom=Side(style='medium', color=NAVY)))
    ws.row_dimensions[row].height = height


def inp(c, fmt=None):
    c.fill = fill(INPUT_FILL)
    c.font = font(10, False, INPUT_FONT)
    c.border = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
    if fmt:
        c.number_format = fmt


def scr(c, fmt=None):
    c.fill = fill(SCRIPT_FILL)
    c.font = font(10, False, '375623')
    if fmt:
        c.number_format = fmt


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def band(ws, r, c1, c2):
    if r % 2 == 0:
        for c in range(c1, c2 + 1):
            if ws.cell(r, c).fill is None or ws.cell(r, c).fill.start_color.rgb in (None, '00000000'):
                ws.cell(r, c).fill = fill(PALE)


def tier_format(ws, rng):
    first = rng.split(':')[0]
    for t, (bg, fg, it) in TIER.items():
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'{first}="{t}"'], fill=fill(bg), font=Font(color=fg, bold=(t == 'Core'), italic=it)))


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def qcol(slot, field):
    return CL(2 + 6 * (slot - 1) + QF.index(field))


def msheet(slot):
    return f'M{slot:02d}'


WRAP = Alignment(wrap_text=True, vertical='center')
WRAPT = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center')
MONEY, PRICE, PCT, PCT1, INT, DATE = '$#,##0', '#,##0.00', '0.0%', '0.0%', '#,##0', 'dd-mmm-yy'


# ----------------------------------------------------------------- Parameters
PARAMS = [
    # (section, name, label, value, unit/format, what, evidence, review)
    ('ACCOUNT  —  written by Script 1 each morning (or typed)', None, None, None, None, None, None, None),
    (None, 'P_Cash', 'Cash available today', 6_900_000, MONEY,
     "Settled cash IBKR reports this morning. Every buy order is sized from it; orders never exceed it.",
     'Cash account (§3.89)', 'Daily · Script'),
    (None, 'P_NetLiq', 'Net liquidation (IBKR)', 6_900_000, MONEY,
     'For reconciliation against the book value the workbook computes. Not used for sizing.',
     '', 'Daily · Script'),
    ('RISK CONTROLS', None, None, None, None, None, None, None),
    (None, 'P_NewBuys', 'New buys allowed', 'YES', '@',
     'Emergency switch. NO = no new buy orders today; targets and stops still run.',
     '', 'Daily'),
    (None, 'P_ExpCap', 'Exposure cap per name', 0.25, PCT,
     "A name's held value plus today's orders may not exceed this share of the book. Applies to "
     'model and discretionary positions alike; the excess is held in cash.',
     'Tested 10–30%; 25% = safety rail, −1.3pp (§3.97)', 'Weekly look · change quarterly'),
    (None, 'P_LiqDefault', 'Liquidity cap — default', 0.10, PCT,
     "A name's orders may not exceed this share of its median first-hour dollar volume. Per-name "
     'overrides (CF 2%) are on the Slots sheet.', 'Tested 1–10% (§3.100)', 'Weekly'),
    (None, 'P_LiqDays', 'Liquidity lookback (sessions)', 20, INT,
     'Sessions in the first-hour volume median.', '§3.100', 'Quarterly'),
    ('BREADTH GATE AND BENCH LENDING', None, None, None, None, None, None, None),
    (None, 'P_GateK', 'Gate threshold (Core names below average)', 6, INT,
     'The gate is ON when at least this many Core names closed below their own average. While it is '
     'on, a Core name below its own average places no buy and its share is held back.',
     'Tested 3–6 of 12 (§3.94); keep at half the Core count', 'Quarterly · two-regime backtest'),
    (None, 'P_GateDMA', 'Gate average length (days)', 200, INT,
     'Length of the moving average the gate compares each Core close with.',
     '50/100 tested and rejected (§3.96)', 'Quarterly · two-regime backtest'),
    (None, 'P_Lend', 'Lend held-back cash to the Bench', 'YES', '@',
     'YES = option E: the gated names\' held-back cash funds the Bench orders (each ≤ one ordinary '
     'order). NO = hold it all in cash.', '§3.90, §3.94, §3.95', 'Quarterly'),
    ('ORDER RULES', None, None, None, None, None, None, None),
    (None, 'P_PMThr', 'Pre-market veto threshold', 0.04, PCT,
     'No buy order when its level sits more than this below the 09:00 pre-market price.',
     '§3.16', 'Quarterly'),
    (None, 'P_ATHeps', 'ATH gate ε', 0.0, '0.0%',
     'Each buy is capped at ATH×(1−ε) − close×premium, so no trade needs a print above the high.',
     '§3.86 (moves results ≤1pp)', 'Quarterly'),
    (None, 'P_StopDays', 'Time stop (calendar days)', 50, INT,
     'A position not at target after this many calendar days is sold at the next open.',
     'Uniform 50 days; per-name stops overfit (§3.81–3.82)', 'Annual'),
    (None, 'P_BayesPct', 'Bayes share of a paper Model sheet', 0.5, PCT,
     'Split of each Model sheet\'s paper capital between its two sleeves. Live orders size each '
     'sleeve equally.', '', 'Annual'),
    ('COSTS AND CAPITAL', None, None, None, None, None, None, None),
    (None, 'P_BuyFee', 'Buy commission ($/share)', 0.005, '0.0000', 'IBKR buy commission per share.', '', 'Monthly'),
    (None, 'P_SellFee', 'Sell commission and fees ($/share)', 0.0095, '0.0000',
     'IBKR sell commission plus regulatory fees per share.', '', 'Monthly'),
    (None, 'P_Interest', 'Interest on idle cash (p.a.)', 0.0314, '0.00%',
     'Rate assumed in the backtests. Confirm IBKR actually pays it on USD balances.',
     '§3.99 (~1.2pp of return)', 'Monthly'),
    (None, 'P_TargetCap', 'Target capital', 10_000_000, MONEY,
     'Planned book size. Liquidity was checked at $10m (§3.99–3.100).', '', 'Monthly'),
    (None, 'P_PaperCap', 'Paper capital per Model sheet', 1_000_000, MONEY,
     'Capital each Model sheet\'s own paper record runs on (Shadow evidence).', '', 'Annual'),
    ('DISCRETIONARY TRADES', None, None, None, None, None, None, None),
    (None, 'P_DiscStop', 'Discretionary time stop (sessions)', 10, INT,
     'A discretionary ticket still open after this many sessions is flagged PAST STOP.',
     '§3.31', 'Monthly'),
    (None, 'P_DiscTarget', 'Discretionary default target', 0.05, PCT,
     'Default take-profit for a DISC ticket, multiplicative off the fill.', '§3.28d', 'Monthly'),
    ('ADMISSION THRESHOLDS', None, None, None, None, None, None, None),
    (None, 'P_BetaMax', 'Max AI beta for a Bench name', 0.27, '0.00',
     'Test 1. At or above this a name counts as AI-related and can only replace a Core name.',
     'G1 screen', 'Quarterly'),
    (None, 'P_RhoPairs', 'Fit test: window pairs needed (of 4)', 3, INT,
     'Test 2. Pairs of 12-month windows with rho ≥ 0.20 needed for fitted settings.',
     '§3.101', 'Quarterly'),
    (None, 'P_ShadowDays', 'Shadow: minimum days', 90, INT,
     'Test 5. Minimum time in Shadow before review.', '§3.99', 'Quarterly'),
    (None, 'P_ShadowFills', 'Shadow: minimum paper fills', 10, INT,
     'Test 5. Minimum paper fills in Shadow before review.', '§3.99', 'Quarterly'),
    (None, 'P_StaleDays', 'Data stale after (calendar days)', 4, INT,
     'A slot whose latest close is older than this is flagged STALE.', '', 'Monthly'),
]


def build_parameters(wb):
    ws = wb.create_sheet('Parameters')
    title(ws, 'PARAMETERS', 'Every rule the book runs on, in one place. Change a blue cell only on its '
          'review cadence, and record the change in the Change Log.')
    widths(ws, {'B': 40, 'C': 14, 'D': 64, 'E': 40, 'F': 24, 'G': 13, 'H': 30})
    r = 4
    for sec, nm, lab, val, fmt, what, ev, rev in PARAMS:
        if sec:
            r += 1
            bar(ws, r, 2, 8, sec)
            r += 1
            header(ws, r, 2, ['Parameter', 'Value', 'What it does', 'Tested / evidence', 'Review',
                              'Last changed', 'Note'], 18)
            r += 1
            continue
        put(ws, f'B{r}', lab, f=font(10, True), al=Alignment(vertical='center', wrap_text=True))
        c = ws[f'C{r}']
        c.value = val
        (scr if nm in ('P_Cash', 'P_NetLiq') else inp)(c, fmt)
        c.alignment = Alignment(horizontal='right', vertical='center')
        put(ws, f'D{r}', what, al=WRAP)
        put(ws, f'E{r}', ev, f=font(9, False, '595959'), al=WRAP)
        put(ws, f'F{r}', rev, f=font(9, True, BLUE), al=WRAP)
        inp(ws[f'G{r}'], DATE)
        ws[f'G{r}'].value = dt.date(2026, 10, 23)
        ws[f'G{r}'].alignment = CENTER
        inp(ws[f'H{r}'])
        ws.row_dimensions[r].height = 34
        name(wb, nm, f'Parameters!$C${r}')
        r += 1
    dvs = [DataValidation(type='list', formula1='"YES,NO"', allow_blank=False)]
    for dv in dvs:
        ws.add_data_validation(dv)
    for row in ws.iter_rows(min_row=5, max_row=r, min_col=3, max_col=3):
        for c in row:
            if c.value in ('YES', 'NO'):
                dvs[0].add(c)
    ws.freeze_panes = 'B4'
    ws.sheet_properties.tabColor = NAVY
    return ws


# ----------------------------------------------------------------- Slots
SHADOW_WHY = {
    'PLTR': 'Settings do not carry over (fit test anti-predictive in 3 of 4 window pairs); −65% in '
            '2022 and lost money as a Bench name in the 2022 test.',
    'TSLA': 'Fit test ≤ 0 in every window pair; average trade 0.8% against a 6.9% spread (≈12 years '
            'to judge); −65% in 2022.',
    'RKLB': 'Removed by risk decision 2 Sep 2026; lost $0.8m in the overbooked D+A test; fit test '
            'unstable (−0.23 to +0.66); −69% in 2022.',
    'APLD': 'Strongest standalone record but a hindsight pick (+538%); listed 2022, so no full 2022 '
            'test; no volume data for the liquidity check.',
    'AAOI': 'Hindsight pick (+628%); 79% of its profit in three months; no 2022 test yet; no volume '
            'data.',
    'CEG': 'AI-related (β 0.75), so it duplicates Core exposure; worsened drawdown as a Bench filler.',
    'UAL': 'Borderline β 0.45; −0.8pp and +1.6pp drawdown as a Bench filler; 2022 file exists for '
           'the book test.',
    'LUNR': 'Helped slightly as a Bench filler (+0.5pp) but a small-cap AI-adjacent multi-bagger; '
            '10% trade spread; listed 2023 (no 2022 test); no volume data.',
}
NOTE = {
    'CF': 'Liquidity cap 2% (§3.100): the tightest name at $10m.',
    'NEM': 'Bench diversifier. No volume column yet: liquidity unmeasured.',
    'LEN': 'Bench diversifier.',
    'DLTR': 'Bench diversifier. No volume column yet: liquidity unmeasured.',
    'RKT': 'Bench diversifier on NEUTRAL settings (fit test fails in all 4 pairs). No volume yet.',
    'TEAM': 'Bench diversifier. No volume column yet: liquidity unmeasured.',
}
VEC = ('lam', 'phi_L', 'psi', 'k', 'premium', 'peak_cap', 'ou_buf_k', 'ou_prem', 'ou_cap', 'ou_W')
VEC_LAB = ('λ', 'φ_L', 'ψ', 'k', 'Bayes π', 'Peak cap', 'OU buffer k', 'OU π', 'OU cap', 'OU W')


def build_slots(wb, seed):
    ws = wb.create_sheet('Slots')
    title(ws, 'SLOTS — the register', 'One row per slot. Tier decides what a slot does; the ticker and its '
          'settings live here. Follow the Guide before changing a Ticker or a Tier.')
    bar(ws, 4, 2, 31, 'SLOT REGISTER')
    put(ws, 'B5', 'Blue = typed. Tier: Core trades · Bench trades only with the gate\'s held-back cash · '
        'Shadow = paper only · Off = holds its cash.', f=font(9, False, '595959', True))
    labels = ['Slot', 'Ticker', 'Tier', 'Group', 'Model sheet', 'AI beta', 'Settings basis',
              'Fit rho (original)', 'Fit pairs ≥0.20 (of 4)', 'Liquidity cap L (blank = default)',
              'Tier since', 'In Shadow since', 'Why not in the book / note',
              'T2 Fit result', 'T3 Book test result', 'T3 evidence', 'T4 Liquidity result',
              'T4 evidence'] + list(VEC_LAB) + ['Core #', 'Bench #', 'Shadow #', 'Off #']
    header(ws, 6, 2, labels, 42)
    since = dt.date(2026, 10, 23)
    for k, s in enumerate(seed['slots'], start=1):
        r = SL0 + k - 1
        t = s['ticker']
        put(ws, f'B{r}', k, fmt='00', al=CENTER)
        put(ws, f'C{r}', t, al=CENTER)
        inp(ws[f'C{r}'])
        put(ws, f'D{r}', s['tier'], al=CENTER)
        inp(ws[f'D{r}'])
        put(ws, f'E{r}', f'=IF(G{r}="","",IF(G{r}>=P_BetaMax,"AI-related","Diversifier"))', al=CENTER)
        put(ws, f'F{r}', msheet(k), al=CENTER, f=font(10, False, '595959'))
        put(ws, f'G{r}', s['beta'], fmt='0.00')
        inp(ws[f'G{r}'], '0.00')
        basis = 'Neutral' if 'NEUTRAL' in s['vector_basis'] else (
            'Fitted' if s['rho'] is not None else 'Live workbook')
        put(ws, f'H{r}', basis)
        inp(ws[f'H{r}'])
        put(ws, f'I{r}', s['rho'] if s['rho'] is not None else '', fmt='+0.00;-0.00')
        inp(ws[f'I{r}'], '+0.00;-0.00')
        w = s['rho_windows']
        pairs = sum(1 for v in w.values() if v is not None and v >= 0.20) if w else ''
        put(ws, f'J{r}', pairs)
        inp(ws[f'J{r}'], '0')
        put(ws, f'K{r}', 0.02 if t == 'CF' else None)
        inp(ws[f'K{r}'], '0.0%')
        put(ws, f'L{r}', since)
        inp(ws[f'L{r}'], DATE)
        put(ws, f'M{r}', since if s['tier'] == 'Shadow' else None)
        inp(ws[f'M{r}'], DATE)
        put(ws, f'N{r}', SHADOW_WHY.get(t, NOTE.get(t, '')), al=WRAP)
        inp(ws[f'N{r}'])
        if s['tier'] == 'Shadow':
            t2 = 'FAIL' if t in ('PLTR', 'TSLA', 'RKLB') else 'NOT RUN'
            vals = [t2, 'FAIL' if t in ('PLTR', 'TSLA', 'RKLB', 'CEG', 'UAL') else 'NOT RUN',
                    {'PLTR': '2022 bench test: lost money (§3.94)', 'TSLA': '2022 bench test: lost money (§3.94)',
                     'RKLB': '2022 bench test: lost money (§3.94)', 'CEG': 'Worsened DD as filler (§3.92)',
                     'UAL': '−0.8pp, +1.6pp DD as filler (§3.92)'}.get(t, 'To run with 2022 data'),
                    'NOT RUN' if t in ('APLD', 'AAOI', 'LUNR') else 'NOT RUN',
                    'No volume column' if t in ('APLD', 'AAOI', 'LUNR') else 'To measure at $10m']
        else:
            vals = ['PASS' if (pairs != '' and pairs >= 3) else ('—' if pairs == '' else 'NEUTRAL'),
                    'PASS', '§3.93–3.95' if s['tier'] == 'Bench' else 'Live book',
                    'PASS' if t not in ('NEM', 'DLTR', 'RKT', 'TEAM', 'CF') else (
                        'CAPPED 2%' if t == 'CF' else 'NOT MEASURED'),
                    '§3.100' if t not in ('NEM', 'DLTR', 'RKT', 'TEAM') else 'Needs volume re-export']
        for j, v in enumerate(vals):
            c = ws[f'{CL(15 + j)}{r}']
            c.value = v
            inp(c)
            c.alignment = WRAP if j in (2, 4) else CENTER
        for j, key in enumerate(VEC):
            c = ws[f'{CL(20 + j)}{r}']
            c.value = round(s['vector'][key]) if key == 'ou_W' else s['vector'][key]
            inp(c, '0' if key == 'ou_W' else '0.000000')
        for j, tr in enumerate(('Core', 'Bench', 'Shadow', 'Off')):
            put(ws, f'{CL(30 + j)}{r}', f'=IF($D{r}="{tr}",COUNTIF($D${SL0}:$D{r},"{tr}"),"")',
                al=CENTER, f=font(9, False, '808080'))
        ws.row_dimensions[r].height = 30 if s['tier'] == 'Shadow' else 18
    last = SL0 + NS - 1
    tier_format(ws, f'D{SL0}:D{last}')
    dv = DataValidation(type='list', formula1='"Core,Bench,Shadow,Off"', allow_blank=False)
    dv.error, dv.errorTitle = 'Choose Core, Bench, Shadow or Off', 'Tier'
    ws.add_data_validation(dv)
    dv.add(f'D{SL0}:D{last}')
    dv2 = DataValidation(type='list', formula1='"PASS,FAIL,NEUTRAL,NOT RUN,NOT MEASURED,CAPPED 2%,—"',
                         allow_blank=True)
    ws.add_data_validation(dv2)
    for col in ('O', 'P', 'R'):
        dv2.add(f'{col}{SL0}:{col}{last}')
    widths(ws, {'B': 5, 'C': 8, 'D': 8, 'E': 11, 'F': 7, 'G': 7, 'H': 11, 'I': 8, 'J': 8, 'K': 10,
                'L': 10, 'M': 10, 'N': 58, 'O': 9, 'P': 9, 'Q': 26, 'R': 14, 'S': 22})
    for j in range(10):
        ws.column_dimensions[CL(20 + j)].width = 9
    for j in range(4):
        ws.column_dimensions[CL(30 + j)].width = 6
    bar(ws, last + 2, 2, 31, 'WHAT THE COLUMNS MEAN')
    notes = [
        'Ticker and Tier are the only cells that change the book. Change them only through a procedure in the Guide, and log it.',
        'Group is derived from AI beta (≥ P_BetaMax = AI-related). An AI-related name can only enter as a Core replacement.',
        'Settings basis: Fitted (fit test passes), Neutral (fit test fails or not run: the middle-of-the-road vector), or Live workbook (the deployed vectors).',
        'T2–T4 results are written after the quarterly re-run of the admission tests; the evidence columns say where they come from.',
        'The ten settings columns (λ … OU W) feed this slot\'s Model sheet. Change them only on annual review with walk-forward evidence (§3.1).',
        'Core # / Bench # / Shadow # / Off # number the slots within each tier; the panels use them to list each tier.',
    ]
    for i, t in enumerate(notes):
        put(ws, f'B{last + 3 + i}', f'{i + 1}.  {t}', f=font(9, False, '404040'))
    ws.freeze_panes = f'D{SL0}'
    ws.sheet_properties.tabColor = '548235'
    name(wb, 'IBKR_Slots', f'Slots!$B${SL0}:$D${last}')
    return ws


# ----------------------------------------------------------------- Query
def build_query(wb, seed):
    ws = wb.create_sheet('Query')
    ws['A1'] = 'Date'
    for k in range(1, NS + 1):
        for f in QF:
            ws[f'{qcol(k, f)}1'] = f'S{k:02d}_{f}'
    hdr = font(9, True, NAVY)
    for c in ws[1]:
        c.font, c.fill = hdr, fill(LIGHT)
    dates = [dt.date.fromisoformat(d) for d in seed['dates']]
    for i, d in enumerate(dates):
        r = Q0 + i
        ws.cell(r, 1, d).number_format = 'yyyy-mm-dd'
        for k, s in enumerate(seed['slots'], start=1):
            row = s['rows'][i]
            if row is None:
                continue
            base = 2 + 6 * (k - 1)
            for j, v in enumerate(row):
                if v is not None:
                    ws.cell(r, base + j, round(v, 6) if j < 4 else round(v))
    ws.freeze_panes = 'B2'
    ws.column_dimensions['A'].width = 11
    ws.sheet_properties.tabColor = '7F7F7F'
    name(wb, 'IBKR_QueryAnchor', 'Query!$A$1')
    return ws, len(dates)


# ----------------------------------------------------------------- Model sheets
def model_templates(live_path):
    """Template formulas (cols A..BF) for rows 8, 9 and the general row 10,
    translated once to every capacity row, Feed refs mapped to Query slot 1."""
    m = openpyxl.load_workbook(live_path)['Model TSM']
    tpl = {r: {c: m.cell(r, c).value for c in range(1, 59)} for r in (1, 2, 3, 6, 7, 8, 9, 10)}
    fmap = {'A': 'A', 'B': qcol(1, 'O'), 'C': qcol(1, 'H'), 'D': qcol(1, 'L'), 'E': qcol(1, 'C')}

    def to_query(f):
        if not isinstance(f, str):
            return f
        return re.sub(r"'Feed TSM'!\$?([A-E])\$?(\d+)",
                      lambda mm: f'Query!{fmap[mm.group(1)]}{mm.group(2)}', f)

    rows = {}
    for r in (8, 9):
        rows[r] = {c: to_query(v) for c, v in tpl[r].items()}
    gen = {c: to_query(v) for c, v in tpl[10].items()}
    for r in range(10, M1 + 1):
        rows[r] = {}
        for c, v in gen.items():
            if isinstance(v, str) and v.startswith('='):
                rows[r][c] = Translator(v, origin=f'{CL(c)}10').translate_formula(f'{CL(c)}{r}')
            else:
                rows[r][c] = v
    return tpl, rows


QREF = re.compile(r'Query!([A-Z]+)(\d+)')


def build_model(wb, slot, tpl, rows):
    ws = wb.create_sheet(msheet(slot))
    sr = SL0 + slot - 1
    cmap = {qcol(1, f): qcol(slot, f) for f in ('O', 'H', 'L', 'C')}
    put(ws, 'A1', f'="Slot {slot:02d} · "&Slots!$C${sr}&" · "&Slots!$D${sr}&'
        f'"  —  Bayes (Kalman) + OU, two sleeves, time stop. Paper record on P_PaperCap."',
        f=font(12, True, NAVY))
    # parameter rows: labels from the template, values from Slots / Parameters
    for c in range(1, 23):
        v = tpl[2].get(c)
        if c % 2 == 1 and v is not None:
            put(ws, f'{CL(c)}2', v, f=font(9, True, '404040'))
    for c in (1, 3, 5, 7, 9):
        v = tpl[3].get(c)
        if v is not None:
            put(ws, f'{CL(c)}3', v, f=font(9, True, '404040'))
    pv = {'B2': f'=Slots!$T${sr}', 'D2': f'=Slots!$U${sr}', 'F2': f'=Slots!$V${sr}',
          'H2': f'=Slots!$W${sr}', 'J2': f'=Slots!$X${sr}', 'L2': f'=Slots!$Y${sr}',
          'N2': '=P_BuyFee', 'P2': '=P_PaperCap', 'R2': '=P_Interest', 'T2': '=P_StopDays',
          'V2': '=P_BayesPct', 'B3': f'=ROUND(Slots!$AC${sr},0)', 'D3': f'=Slots!$Z${sr}',
          'H3': f'=Slots!$AA${sr}', 'J3': f'=Slots!$AB${sr}',
          'F3': f'=IFERROR(CORREL(OFFSET($BE$8,$B$3+1,0,{N - 1}-$B$3,1),'
                f'OFFSET($BF$8,$B$3+1,0,{N - 1}-$B$3,1)),"")'}
    for ref, v in pv.items():
        put(ws, ref, v, f=font(10, False, '375623'), fl=SCRIPT_FILL)
    stats = {'X4': 'Total profit:', 'Y4': f'=Y{M1}+AF{M1}-$P$2', 'Z4': 'Total buys:',
             'AA4': f'=SUM(Z{M0}:Z{M1})+SUM(AG{M0}:AG{M1})', 'AB4': 'Stop-loss exits:',
             'AC4': f'=SUMPRODUCT((AD{M0}:AD{M1}=1)*(AC{M0}:AC{M1}<AB{M0}:AB{M1}))'
                    f'+SUMPRODUCT((AK{M0}:AK{M1}=1)*(AJ{M0}:AJ{M1}<AI{M0}:AI{M1}))',
             'X5': 'Annual return:',
             'Y5': f'=IFERROR(((Y{M1}+AF{M1})/$P$2)^(365.25/(MAX($A${M0}:$A${M1})-MIN($A${M0}:$A${M1})))-1,"")',
             'Z5': 'Total interest:', 'AA5': f'=SUM(AQ{M0}:AQ{M1})+SUM(AS{M0}:AS{M1})'}
    for ref, v in stats.items():
        put(ws, ref, v, f=font(9, isinstance(v, str) and not v.startswith('='), '404040'),
            fmt=MONEY if ref in ('Y4', 'AA5') else ('0.0%' if ref == 'Y5' else None))
    for c in range(1, 59):
        for rr in (6, 7):
            v = tpl[rr].get(c)
            if v is not None:
                put(ws, f'{CL(c)}{rr}', v, f=font(9, True, NAVY), fl=LIGHT if rr == 7 else None)
    for r in range(M0, M1 + 1):
        for c, v in rows[r].items():
            if v is None:
                continue
            if c <= 5 and isinstance(v, str):
                v = QREF.sub(lambda mm: f'Query!{cmap.get(mm.group(1), mm.group(1))}{mm.group(2)}', v)
            ws.cell(r, c, v)
    ws.freeze_panes = 'F8'
    ws.column_dimensions['A'].width = 10
    ws.sheet_properties.tabColor = 'A6A6A6'
    for c in (1,):
        for r in range(M0, M1 + 1):
            ws.cell(r, c).number_format = 'yyyy-mm-dd'
    return ws


# ----------------------------------------------------------------- Engine
def build_engine(wb):
    ws = wb.create_sheet('Engine')
    title(ws, 'ENGINE — levels, gate and order sizing',
          'Calculated, not edited. Levels per slot (next-session bids), the breadth gate, then each sleeve\'s '
          'order: base share → Bench lending → exposure and liquidity caps → shares.')
    widths(ws, {'B': 44, 'C': 16})
    bar(ws, 3, 2, 3, 'BOOK STATE AND SIZING')
    lv1, sv1 = LV0 + NS - 1, SV0 + 2 * NS - 1
    g = [
        ('Equity — cash + open positions at last close', '=P_Cash+SUM(Log!$T$%d:$T$%d)' % (LG0, LG1), MONEY, 'E_Equity'),
        ('Core names', f'=COUNTIF($C${LV0}:$C${lv1},"Core")', INT, 'E_CoreN'),
        ('Core names below their own average', f'=COUNTIFS($C${LV0}:$C${lv1},"Core",$AC${LV0}:$AC${lv1},TRUE)', INT, 'E_Below'),
        ('Breadth gate ON?', '=E_Below>=P_GateK', None, 'E_GateOn'),
        ('Sleeves bidding — Core', f'=COUNTIF($L${SV0}:$L${sv1},"core")', INT, 'E_nCore'),
        ('Sleeves gated — share held back', f'=COUNTIF($L${SV0}:$L${sv1},"gated")', INT, 'E_nGated'),
        ('Sleeves Off — share held in cash', f'=COUNTIF($L${SV0}:$L${sv1},"off")', INT, 'E_nOff'),
        ('Sleeves bidding — Bench', f'=COUNTIF($L${SV0}:$L${sv1},"bench")', INT, 'E_nBench'),
        ('Shares in the morning split', '=E_nCore+E_nGated+E_nOff', INT, 'E_Den'),
        ('Ordinary order per sleeve ($)', '=IF(E_Den=0,0,P_Cash/E_Den)', MONEY, 'E_Share'),
        ('Held-back cash available to lend ($)', '=IF(AND(P_Lend="YES",E_Den>0),P_Cash*E_nGated/E_Den,0)', MONEY, 'E_Lend'),
        ('Bench order per sleeve ($)', '=IF(E_nBench=0,0,MIN(E_Lend/E_nBench,E_Share))', MONEY, 'E_BenchEach'),
        ('Total buy orders today ($)', f'=SUM($Q${SV0}:$Q${sv1})', MONEY, 'E_Orders'),
        ('Lent to the Bench today ($)', f'=SUMIFS($Q${SV0}:$Q${sv1},$L${SV0}:$L${sv1},"bench")', MONEY, 'E_Lent'),
        ('Cash not ordered today ($)', '=P_Cash-E_Orders', MONEY, 'E_Idle'),
    ]
    for i, (lab, f, fmt, nm) in enumerate(g):
        r = 4 + i
        put(ws, f'B{r}', lab)
        put(ws, f'C{r}', f, fmt=fmt, f=font(10, True))
        name(wb, nm, f'Engine!$C${r}')
    # ---- levels
    bar(ws, LV0 - 2, 2, 47, 'LEVELS PER SLOT — next-session bids, gate, caps')
    lab = ['Slot', 'Ticker', 'Tier', 'Model', 'Data rows', 'Last pos', 'Last date', 'Last close',
           'Bayes fair', 'Sigma', 'ATH', 'k', 'Bayes π', 'Peak cap', 'Bayes buy', 'Bayes target',
           'OU W', 'OU mean', 'OU AR', 'OU sigma', 'OU forecast', 'OU buffer k', 'OU π', 'OU cap',
           'OU buy', 'OU target', 'Gate average', 'vs average', 'Below?', 'Gated?', 'FH pos',
           'FH as of', 'First-hour $ median', 'Liquidity L', 'Liquidity cap $', 'Held value $',
           'Exposure %', 'Exposure room $', 'Orders pre-cap $', 'Name cap $', 'Scale', 'Binding cap',
           'Data', 'Live fills (28d)', 'Model fills (20 sess.)', 'Shadow since', 'Paper fills since',
           'Paper return since', 'Paper stop-outs since']
    header(ws, LV0 - 1, 1, lab, 42)
    for k in range(1, NS + 1):
        r = LV0 + k - 1
        sr = SL0 + k - 1
        m = msheet(k)
        E = f"{m}!$E${M0}:$E${M1}"
        fh = f"Query!${qcol(k, 'FH')}${Q0}:${qcol(k, 'FH')}${Q1}"
        f = {
            'A': k, 'B': f'=Slots!$C${sr}', 'C': f'=Slots!$D${sr}', 'D': m,
            'E': f'=COUNT({E})', 'F': f'=IF(E{r}=0,"",MATCH(9.99E+307,{E}))',
            'G': f'=IF(F{r}="","",INDEX({m}!$A${M0}:$A${M1},F{r}))',
            'H': f'=IF(F{r}="","",INDEX({E},F{r}))',
            'I': f'={m}!R{M1}+{m}!S{M1}', 'J': f'={m}!W{M1}', 'K': f'={m}!G{M1}',
            'L': f'={m}!$H$2', 'M': f'={m}!$J$2', 'N': f'={m}!$L$2',
            'O': f'=IF(H{r}="","",MIN(I{r}-L{r}*J{r},K{r}*(1-N{r}),K{r}*(1-P_ATHeps)-H{r}*M{r}))',
            'P': f'=IF(O{r}="","",O{r}+H{r}*M{r})',
            'Q': f'={m}!$B$3',
            'R': f'=IF(OR(F{r}="",F{r}<Q{r}),"",AVERAGE(INDEX({E},F{r}-Q{r}+1):INDEX({E},F{r})))',
            'S': f'=IF(R{r}="","",MEDIAN(0,SLOPE(INDEX({E},F{r}-Q{r}+2):INDEX({E},F{r}),'
                 f'INDEX({E},F{r}-Q{r}+1):INDEX({E},F{r}-1)),0.99))',
            'T': f'=IF(R{r}="","",SQRT(SUMPRODUCT((((INDEX({E},F{r}-Q{r}+2):INDEX({E},F{r})-R{r})'
                 f'-S{r}*(INDEX({E},F{r}-Q{r}+1):INDEX({E},F{r}-1)-R{r}))'
                 f'-((AVERAGE(INDEX({E},F{r}-Q{r}+2):INDEX({E},F{r}))-R{r})'
                 f'-S{r}*(AVERAGE(INDEX({E},F{r}-Q{r}+1):INDEX({E},F{r}-1))-R{r})))^2)/(Q{r}-1)))',
            'U': f'=IF(R{r}="","",R{r}+S{r}*(H{r}-R{r}))',
            'V': f'={m}!$D$3', 'W': f'={m}!$H$3', 'X': f'={m}!$J$3',
            'Y': f'=IF(U{r}="","",MIN(U{r}-V{r}*T{r},K{r}*(1-X{r}),K{r}*(1-P_ATHeps)-H{r}*W{r}))',
            'Z': f'=IF(Y{r}="","",Y{r}+H{r}*W{r})',
            'AA': f'=IF(OR(F{r}="",F{r}<P_GateDMA),"",AVERAGE(INDEX({E},F{r}-P_GateDMA+1):INDEX({E},F{r})))',
            'AB': f'=IF(AA{r}="","",H{r}/AA{r}-1)',
            'AC': f'=IF(AA{r}="",FALSE,H{r}<AA{r})',
            'AD': f'=AND(C{r}="Core",AC{r},E_GateOn)',
            'AE': f'=IFERROR(MATCH(9.99E+307,{fh}),"")',
            'AF': f'=IF(AE{r}="","",INDEX(Query!$A${Q0}:$A${Q1},AE{r}))',
            'AG': f'=IF(AE{r}="","",MEDIAN(INDEX({fh},MAX(1,AE{r}-P_LiqDays+1)):INDEX({fh},AE{r})))',
            'AH': f'=IF(Slots!$K${sr}="",P_LiqDefault,Slots!$K${sr})',
            'AI': f'=IF(AG{r}="","",AH{r}*AG{r})',
            'AJ': f'=SUMIFS(Log!$T${LG0}:$T${LG1},Log!$C${LG0}:$C${LG1},B{r})',
            'AK': f'=IF(E_Equity>0,AJ{r}/E_Equity,0)',
            'AL': f'=MAX(0,P_ExpCap*E_Equity-AJ{r})',
            'AM': f'=SUMIFS($M${SV0}:$M${SV0 + 2 * NS - 1},$A${SV0}:$A${SV0 + 2 * NS - 1},A{r})',
            'AN': f'=IF(AI{r}="",AL{r},MIN(AL{r},AI{r}))',
            'AO': f'=IF(AM{r}<=0,1,MIN(1,AN{r}/AM{r}))',
            'AP': f'=IF(AO{r}>=1,"",IF(OR(AI{r}="",AL{r}<=AI{r}),"Exposure","Liquidity"))',
            'AQ': f'=IF(G{r}="","NO DATA",IF(TODAY()-G{r}>P_StaleDays,"STALE","OK"))',
            'AR': f'=COUNTIFS(Log!$C${LG0}:$C${LG1},B{r},Log!$B${LG0}:$B${LG1},"Model",Log!$E${LG0}:$E${LG1},">="&(TODAY()-28))',
            'AS': f'=IF(F{r}="","",SUM(INDEX({m}!$Z${M0}:$Z${M1},MAX(1,F{r}-19)):INDEX({m}!$Z${M0}:$Z${M1},F{r}))'
                  f'+SUM(INDEX({m}!$AG${M0}:$AG${M1},MAX(1,F{r}-19)):INDEX({m}!$AG${M0}:$AG${M1},F{r})))',
            'AT': f'=IF(Slots!$M${sr}="","",Slots!$M${sr})',
            'AU': f'=IF(AT{r}="","",SUMPRODUCT(ISNUMBER({m}!$A${M0}:$A${M1})*({m}!$A${M0}:$A${M1}>=AT{r})'
                  f'*({m}!$Z${M0}:$Z${M1}+{m}!$AG${M0}:$AG${M1})))',
            'AV': f'=IF(OR(AT{r}="",F{r}=""),"",IFERROR((INDEX({m}!$BC${M0}:$BC${M1},F{r})+INDEX({m}!$BD${M0}:$BD${M1},F{r}))'
                  f'/(INDEX({m}!$BC${M0}:$BC${M1},MATCH(AT{r},{m}!$A${M0}:$A${M1},1))'
                  f'+INDEX({m}!$BD${M0}:$BD${M1},MATCH(AT{r},{m}!$A${M0}:$A${M1},1)))-1,0))',
            'AW': f'=IF(AT{r}="","",SUMPRODUCT(ISNUMBER({m}!$A${M0}:$A${M1})*({m}!$A${M0}:$A${M1}>=AT{r})'
                  f'*((({m}!$AD${M0}:$AD${M1}=1)*({m}!$AC${M0}:$AC${M1}<{m}!$AB${M0}:$AB${M1}))'
                  f'+(({m}!$AK${M0}:$AK${M1}=1)*({m}!$AJ${M0}:$AJ${M1}<{m}!$AI${M0}:$AI${M1})))))',
        }
        fmts = {'G': DATE, 'H': PRICE, 'I': PRICE, 'J': PRICE, 'K': PRICE, 'O': PRICE, 'P': PRICE,
                'R': PRICE, 'T': PRICE, 'U': PRICE, 'Y': PRICE, 'Z': PRICE, 'AA': PRICE, 'AB': PCT,
                'AF': DATE, 'AG': MONEY, 'AH': PCT, 'AI': MONEY, 'AJ': MONEY, 'AK': PCT, 'AL': MONEY,
                'AM': MONEY, 'AN': MONEY, 'AO': '0.00', 'AT': DATE, 'AV': PCT, 'M': '0.00%', 'N': '0.00%',
                'W': '0.00%', 'X': '0.00%', 'L': '0.000', 'V': '0.000', 'S': '0.000'}
        for col, v in f.items():
            put(ws, f'{col}{r}', v, fmt=fmts.get(col), f=font(9))
    tier_format(ws, f'C{LV0}:C{lv1}')
    # ---- sleeves
    bar(ws, SV0 - 2, 2, 24, 'SLEEVES — two per slot (Bayes, OU); one order each')
    lab2 = ['Slot', 'Ticker', 'Tier', 'Sleeve', 'Model level', 'Buy used', 'Target used', 'Held?',
            'Has bid', '9:00 price', 'PM veto', 'Category', 'Base $', 'Scale', 'Final $', 'Shares',
            'Order $', 'Held shares', 'Held since', 'Held target', 'Days held', 'Stop due?']
    header(ws, SV0 - 1, 1, lab2, 30)
    for k in range(1, NS + 1):
        lr = LV0 + k - 1
        for j, sl in enumerate(('Bayes', 'OU')):
            r = SV0 + 2 * (k - 1) + j
            o = OR0 + 2 * (k - 1) + j
            lvl, pi = ('O', 'M') if sl == 'Bayes' else ('Y', 'W')
            held = (f'Log!$C${LG0}:$C${LG1},B{r},Log!$D${LG0}:$D${LG1},D{r},'
                    f'Log!$P${LG0}:$P${LG1},"OPEN",Log!$B${LG0}:$B${LG1},"Model"')
            f = {
                'A': k, 'B': f'=B{lr}', 'C': f'=C{lr}', 'D': sl,
                'E': f'={lvl}{lr}',
                'F': f'=IF(Orders!$N${o}<>"",Orders!$N${o},E{r})',
                'G': f'=IF(Orders!$O${o}<>"",Orders!$O${o},IF(OR(F{r}="",$H${lr}=""),"",F{r}+$H${lr}*${pi}${lr}))',
                'H': f'=COUNTIFS({held})>0',
                'I': f'=AND(ISNUMBER(F{r}),N(F{r})>0)',
                'J': f'=IF(Orders!$M${o}="","",Orders!$M${o})',
                'K': f'=AND(ISNUMBER(J{r}),I{r},N(F{r})<N(J{r})*(1-P_PMThr))',
                'L': f'=IF(H{r},"held",IF(NOT(I{r}),"no bid",IF(C{r}="Shadow","shadow",IF(C{r}="Off","off",'
                     f'IF(P_NewBuys<>"YES","paused",IF(K{r},"veto",IF(C{r}="Bench","bench",'
                     f'IF($AD${lr},"gated","core"))))))))',
                'M': f'=IF(L{r}="core",E_Share,IF(L{r}="bench",E_BenchEach,0))',
                'N': f'=$AO${lr}',
                'O': f'=M{r}*N{r}',
                'P': f'=IF(AND(O{r}>0,I{r}),FLOOR(O{r}/(F{r}+P_BuyFee),1),0)',
                'Q': f'=P{r}*IF(I{r},F{r}+P_BuyFee,0)',
                'R': f'=SUMIFS(Log!$G${LG0}:$G${LG1},{held})',
                'S': f'=IF(H{r},SUMIFS(Log!$E${LG0}:$E${LG1},{held}),"")',
                'T': f'=IF(H{r},SUMIFS(Log!$I${LG0}:$I${LG1},{held}),"")',
                'U': f'=IF(H{r},TODAY()-S{r},"")',
                'V': f'=IF(H{r},U{r}>=P_StopDays,FALSE)',
            }
            fm = {'E': PRICE, 'F': PRICE, 'G': PRICE, 'J': PRICE, 'M': MONEY, 'N': '0.00', 'O': MONEY,
                  'P': INT, 'Q': MONEY, 'R': INT, 'S': DATE, 'T': PRICE}
            for col, v in f.items():
                put(ws, f'{col}{r}', v, fmt=fm.get(col), f=font(9))
    tier_format(ws, f'C{SV0}:C{sv1}')
    ws.freeze_panes = f'D{LV0}'
    ws.sheet_properties.tabColor = 'A6A6A6'
    return ws


# ----------------------------------------------------------------- Orders
def build_orders(wb):
    ws = wb.create_sheet('Orders')
    title(ws, "TODAY'S ORDERS", 'Script 1 sends every row with Send = YES as a bracket: BUY limit at "Buy @" '
          'with a take-profit SELL limit at "Sell @"; STOP rows sell at the open. Blue cells are yours.')
    widths(ws, {'B': 6, 'C': 8, 'D': 8, 'E': 7, 'F': 10, 'G': 26, 'H': 10, 'I': 10, 'J': 9, 'K': 12,
                'L': 7, 'M': 11, 'N': 11, 'O': 11, 'P': 10, 'Q': 8, 'R': 30})
    kp = [('Data as of', '=MAX(Query!$A:$A)', DATE), ('Cash available', '=P_Cash', MONEY),
          ('Buy orders', '=COUNTIF($G$%d:$G$%d,"BUY*")' % (OR0, OR0 + 2 * NS - 1), INT),
          ('Buy orders $', '=E_Orders', MONEY),
          ('Breadth gate', '=IF(E_GateOn,"ON — "&E_Below&" of "&E_CoreN,"off — "&E_Below&" of "&E_CoreN)', None),
          ('Lent to Bench', '=E_Lent', MONEY)]
    for i, (lab, f, fmt) in enumerate(kp):
        c = 2 + 3 * i
        for rr in (4, 5):
            ws.merge_cells(f'{CL(c)}{rr}:{CL(c + 2)}{rr}')
        put(ws, f'{CL(c)}4', lab, f=font(9, True, '595959'))
        put(ws, f'{CL(c)}5', f, f=font(12, True, NAVY), fmt=fmt, al=Alignment(horizontal='left'))
    put(ws, 'B7', 'A sleeve whose action is not BUY or STOP sends nothing. GATED and OFF keep their cash; '
        'SHADOW shows the paper level only.', f=font(9, False, '595959', True))
    lab = ['Slot', 'Ticker', 'Tier', 'Sleeve', 'Position', 'Action', 'Buy @', 'Sell @', 'Shares',
           'Order $', 'Send?', '9:00 pre-market', 'Buy override', 'Sell override', 'Bought', 'Days',
           'Note']
    header(ws, OR0 - 1, 2, lab, 30)
    for k in range(1, NS + 1):
        lr = LV0 + k - 1
        for j, sl in enumerate(('Bayes', 'OU')):
            r = OR0 + 2 * (k - 1) + j
            s = SV0 + 2 * (k - 1) + j
            act = (f'=IF(F{r}="HOLDING",IF(Engine!$V${s},"STOP DUE — SELL AT OPEN","HOLD — target resting"),'
                   f'IF(Engine!$L${s}="core",IF(J{r}>0,"BUY","CAPPED — no room"),'
                   f'IF(Engine!$L${s}="bench",IF(J{r}>0,"BUY — Bench (lent cash)",'
                   f'IF(E_GateOn,"CAPPED — no room","BENCH — idle (gate off)")),'
                   f'IF(Engine!$L${s}="gated","GATED — cash held",IF(Engine!$L${s}="off","OFF — cash held",'
                   f'IF(Engine!$L${s}="shadow","SHADOW — paper only",IF(Engine!$L${s}="veto","PM VETO — no order",'
                   f'IF(Engine!$L${s}="paused","NEW BUYS PAUSED","NO BID"))))))))')
            f = {
                'B': k, 'C': f'=Engine!$B${lr}', 'D': f'=Engine!$C${lr}', 'E': sl,
                'F': f'=IF(Engine!$H${s},"HOLDING","CASH")',
                'G': act,
                'H': f'=IF(F{r}="HOLDING","",IF(Engine!$I${s},Engine!$F${s},""))',
                'I': f'=IF(F{r}="HOLDING",IF(O{r}<>"",O{r},Engine!$T${s}),IF(H{r}="","",Engine!$G${s}))',
                'J': f'=IF(F{r}="HOLDING",Engine!$R${s},Engine!$P${s})',
                'K': f'=IF(F{r}="HOLDING","",Engine!$Q${s})',
                'L': f'=IF(OR(LEFT(G{r},3)="BUY",LEFT(G{r},4)="STOP"),"YES","NO")',
                'P': f'=IF(F{r}="HOLDING",Engine!$S${s},"")',
                'Q': f'=IF(F{r}="HOLDING",Engine!$U${s},"")',
                'R': f'=IF(AND(F{r}="CASH",Engine!$N${s}<1,Engine!$M${s}>0),"Cut by "&Engine!$AP${lr}&" cap","")',
            }
            fm = {'H': PRICE, 'I': PRICE, 'J': INT, 'K': MONEY, 'P': DATE}
            for col, v in f.items():
                put(ws, f'{col}{r}', v, fmt=fm.get(col), al=CENTER if col in 'BDEFL' else None)
            for col in ('M', 'N', 'O'):
                inp(ws[f'{col}{r}'], PRICE)
            if j == 1:
                for c in range(2, 19):
                    ws.cell(r, c).border = Border(bottom=THIN)
    lo, hi = OR0, OR0 + 2 * NS - 1
    tier_format(ws, f'D{lo}:D{hi}')
    ws.conditional_formatting.add(f'G{lo}:G{hi}', FormulaRule(formula=[f'LEFT(G{lo},3)="BUY"'],
                                  fill=fill(GRN_F), font=Font(color=GRN_T, bold=True)))
    ws.conditional_formatting.add(f'G{lo}:G{hi}', FormulaRule(formula=[f'LEFT(G{lo},4)="STOP"'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add(f'G{lo}:G{hi}', FormulaRule(
        formula=[f'OR(LEFT(G{lo},5)="GATED",LEFT(G{lo},3)="OFF",LEFT(G{lo},2)="PM",LEFT(G{lo},6)="CAPPED")'],
        fill=fill(AMB_F), font=Font(color=AMB_T)))
    ws.conditional_formatting.add(f'G{lo}:G{hi}', FormulaRule(formula=[f'LEFT(G{lo},6)="SHADOW"'],
                                  font=Font(color='808080', italic=True)))
    ws.conditional_formatting.add(f'L{lo}:L{hi}', CellIsRule(operator='equal', formula=['"YES"'],
                                  fill=fill(GRN_F), font=Font(color=GRN_T, bold=True)))
    ws.freeze_panes = f'F{OR0}'
    ws.sheet_properties.tabColor = '548235'
    name(wb, 'IBKR_Orders', f'Orders!$B${OR0 - 1}:$R${hi}')
    name(wb, 'IBKR_BuyFee', 'P_BuyFee')
    name(wb, 'IBKR_SellFee', 'P_SellFee')
    name(wb, 'IBKR_AvailFunds', 'P_Cash')
    return ws


# ----------------------------------------------------------------- Log
def build_log(wb):
    ws = wb.create_sheet('Log')
    title(ws, 'TRADE LOG', 'Every fill, model and discretionary. Script 2 writes the blue columns; you can '
          'type them too. Discretionary tickets count against the same 25% exposure cap.')
    lab = ['#', 'Type', 'Ticker', 'Sleeve', 'Buy date', 'Buy price', 'Shares', 'Buy cost', 'Target',
           'Stop due', 'Sell date', 'Sell price', 'Proceeds', 'Net P&L', 'Return', 'Status',
           'Exit reason', 'Days held', 'Last close', 'Open value', 'Unrealised', 'Flag']
    header(ws, LG0 - 1, 1, lab, 30)
    put(ws, 'B4', 'Type: Model or Disc · Sleeve: Bayes, OU or DISC · Exit reason: Target, Stop or Manual.',
        f=font(9, False, '595959', True))
    widths(ws, {'A': 5, 'B': 7, 'C': 8, 'D': 7, 'E': 10, 'F': 10, 'G': 8, 'H': 12, 'I': 10, 'J': 10,
                'K': 10, 'L': 10, 'M': 12, 'N': 11, 'O': 8, 'P': 8, 'Q': 9, 'R': 6, 'S': 10, 'T': 12,
                'U': 11, 'V': 12})
    lv1 = LV0 + NS - 1
    for r in range(LG0, LG1 + 1):
        f = {
            'A': f'=IF(C{r}="","",ROW()-{LG0 - 1})',
            'H': f'=IF(OR(F{r}="",G{r}=""),"",G{r}*F{r}+G{r}*P_BuyFee)',
            'J': f'=IF(E{r}="","",IF(B{r}="Disc",WORKDAY(E{r},P_DiscStop),E{r}+P_StopDays))',
            'M': f'=IF(OR(L{r}="",G{r}=""),"",G{r}*L{r}-G{r}*P_SellFee)',
            'N': f'=IF(M{r}="","",M{r}-H{r})',
            'O': f'=IF(N{r}="","",N{r}/H{r})',
            'P': f'=IF(E{r}="","",IF(K{r}="","OPEN","CLOSED"))',
            'R': f'=IF(E{r}="","",IF(K{r}="",TODAY(),K{r})-E{r})',
            'S': f'=IF(P{r}="OPEN",IFERROR(INDEX(Engine!$H${LV0}:$H${lv1},MATCH(C{r},Engine!$B${LV0}:$B${lv1},0)),F{r}),"")',
            'T': f'=IF(P{r}="OPEN",G{r}*S{r},0)',
            'U': f'=IF(P{r}="OPEN",T{r}-H{r},"")',
            'V': f'=IF(P{r}<>"OPEN","",IF(TODAY()>=J{r},"PAST STOP",IF(IFERROR(INDEX(Engine!$AK${LV0}:$AK${lv1},'
                 f'MATCH(C{r},Engine!$B${LV0}:$B${lv1},0)),0)>P_ExpCap,"OVER CAP","")))',
        }
        for col, v in f.items():
            ws[f'{col}{r}'] = v
        for col in 'BCDEFGIKLQ':
            c = ws[f'{col}{r}']
            c.fill = fill(INPUT_FILL)
            c.font = font(10, False, INPUT_FONT)
        for col, fmt in (('E', DATE), ('F', PRICE), ('G', INT), ('H', MONEY), ('I', PRICE), ('J', DATE),
                         ('K', DATE), ('L', PRICE), ('M', MONEY), ('N', '$#,##0;[Red]-$#,##0'),
                         ('O', '0.0%;[Red]-0.0%'), ('S', PRICE), ('T', MONEY), ('U', '$#,##0;[Red]-$#,##0')):
            ws[f'{col}{r}'].number_format = fmt
    dv = DataValidation(type='list', formula1='"Model,Disc"', allow_blank=True)
    dv2 = DataValidation(type='list', formula1='"Bayes,OU,DISC"', allow_blank=True)
    dv3 = DataValidation(type='list', formula1='"Target,Stop,Manual"', allow_blank=True)
    for d_, col in ((dv, 'B'), (dv2, 'D'), (dv3, 'Q')):
        ws.add_data_validation(d_)
        d_.add(f'{col}{LG0}:{col}{LG1}')
    ws.conditional_formatting.add(f'V{LG0}:V{LG1}', CellIsRule(operator='notEqual', formula=['""'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add(f'P{LG0}:P{LG1}', CellIsRule(operator='equal', formula=['"OPEN"'],
                                  font=Font(color=BLUE, bold=True)))
    ws.freeze_panes = f'E{LG0}'
    ws.sheet_properties.tabColor = '548235'
    name(wb, 'IBKR_LogAnchor', f'Log!$A${LG0 - 1}')
    name(wb, 'IBKR_Log', f'Log!$A${LG0 - 1}:$V${LG1}')
    return ws


# ----------------------------------------------------------------- Performance
def build_performance(wb):
    ws = wb.create_sheet('Performance')
    title(ws, 'PERFORMANCE', 'Fresh start. Script 2 (or you) adds one NAV row per session; returns are '
          'time-weighted, so deposits and withdrawals do not count as performance.')
    widths(ws, {'A': 11, 'B': 30, 'C': 15, 'E': 6, 'F': 8, 'G': 8, 'H': 9, 'I': 13, 'J': 8, 'K': 8, 'L': 13,
                'N': 11, 'O': 13, 'P': 13, 'Q': 8})
    bar(ws, 4, 2, 3, 'SINCE START')
    nav = f'$B${NAV0}:$B${NAV1}'
    s = [('Start date', f'=IF(COUNT($A${NAV0}:$A${NAV1})=0,"",MIN($A${NAV0}:$A${NAV1}))', DATE),
         ('Latest NAV', f'=IF(COUNT({nav})=0,"",LOOKUP(2,1/({nav}<>""),{nav}))', MONEY),
         ('Return since start (time-weighted)', f'=IF(COUNT($E${NAV0}:$E${NAV1})=0,"",LOOKUP(2,1/($E${NAV0}:$E${NAV1}<>""),$E${NAV0}:$E${NAV1})-1)', PCT),
         ('Max drawdown since start', f'=IF(COUNT($G${NAV0}:$G${NAV1})=0,"",MAX($G${NAV0}:$G${NAV1}))', PCT),
         ('Drawdown — now', f'=IF(COUNT($G${NAV0}:$G${NAV1})=0,"",LOOKUP(2,1/($G${NAV0}:$G${NAV1}<>""),$G${NAV0}:$G${NAV1}))', PCT),
         ('Peak NAV', f'=IF(COUNT({nav})=0,"",MAX({nav}))', MONEY),
         ('Realised P&L — model', f'=SUMIFS(Log!$N${LG0}:$N${LG1},Log!$B${LG0}:$B${LG1},"Model")', '$#,##0;[Red]-$#,##0'),
         ('Realised P&L — discretionary', f'=SUMIFS(Log!$N${LG0}:$N${LG1},Log!$B${LG0}:$B${LG1},"Disc")', '$#,##0;[Red]-$#,##0'),
         ('Unrealised on open positions', f'=SUM(Log!$U${LG0}:$U${LG1})', '$#,##0;[Red]-$#,##0'),
         ('Closed trades', f'=COUNTIF(Log!$P${LG0}:$P${LG1},"CLOSED")', INT),
         ('Win rate', f'=IFERROR(COUNTIFS(Log!$P${LG0}:$P${LG1},"CLOSED",Log!$N${LG0}:$N${LG1},">0")/C14,"")', PCT),
         ('Stop exits', f'=COUNTIF(Log!$Q${LG0}:$Q${LG1},"Stop")', INT)]
    for i, (lab, f, fmt) in enumerate(s):
        r = 5 + i
        put(ws, f'B{r}', lab)
        put(ws, f'C{r}', f, fmt=fmt, f=font(10, True))
    name(wb, 'PF_PeakNAV', 'Performance!$C$10')
    name(wb, 'PF_MaxDD', 'Performance!$C$8')
    # per slot
    bar(ws, 4, 5, 12, 'BY SLOT')
    header(ws, 5, 5, ['Slot', 'Ticker', 'Tier', 'Closed', 'Realised P&L', 'Win %', 'Open', 'Unrealised'], 30)
    for k in range(1, NS + 1):
        r = 5 + k
        sr = SL0 + k - 1
        rng = f'Log!$C${LG0}:$C${LG1},$F{r}'
        f = {'E': k, 'F': f'=Slots!$C${sr}', 'G': f'=Slots!$D${sr}',
             'H': f'=COUNTIFS({rng},Log!$P${LG0}:$P${LG1},"CLOSED")',
             'I': f'=SUMIFS(Log!$N${LG0}:$N${LG1},{rng})',
             'J': f'=IFERROR(COUNTIFS({rng},Log!$P${LG0}:$P${LG1},"CLOSED",Log!$N${LG0}:$N${LG1},">0")/H{r},"")',
             'K': f'=COUNTIFS({rng},Log!$P${LG0}:$P${LG1},"OPEN")',
             'L': f'=SUMIFS(Log!$U${LG0}:$U${LG1},{rng})'}
        for col, v in f.items():
            put(ws, f'{col}{r}', v, fmt={'I': '$#,##0;[Red]-$#,##0', 'J': PCT, 'L': '$#,##0;[Red]-$#,##0'}.get(col),
                al=CENTER if col in 'EFGHK' else None)
    tier_format(ws, f'G6:G{5 + NS}')
    # monthly
    bar(ws, 4, 14, 17, 'BY MONTH (realised, by sell date)')
    header(ws, 5, 14, ['Month', 'Model', 'Discretionary', 'Trades'], 30)
    for i in range(36):
        r = 6 + i
        put(ws, f'N{r}', dt.date(2026, 11, 1) if i == 0 else f'=EDATE(N{r - 1},1)', fmt='mmm-yy')
        win = (f'Log!$K${LG0}:$K${LG1},">="&$N{r},Log!$K${LG0}:$K${LG1},"<"&EDATE($N{r},1)')
        put(ws, f'O{r}', f'=SUMIFS(Log!$N${LG0}:$N${LG1},Log!$B${LG0}:$B${LG1},"Model",{win})', fmt='$#,##0;[Red]-$#,##0')
        put(ws, f'P{r}', f'=SUMIFS(Log!$N${LG0}:$N${LG1},Log!$B${LG0}:$B${LG1},"Disc",{win})', fmt='$#,##0;[Red]-$#,##0')
        put(ws, f'Q{r}', f'=COUNTIFS({win})', fmt=INT)
    # NAV log
    bar(ws, NAV0 - 2, 1, 7, 'DAILY NAV LOG  —  one row per session (Script 2 appends; blue = typed)')
    header(ws, NAV0 - 1, 1, ['Date', 'Net liquidation', 'Deposit (+) / withdrawal (−)', 'Daily return',
                             'Index', 'Peak index', 'Drawdown'], 30)
    for r in range(NAV0, NAV1 + 1):
        first = r == NAV0
        f = {'D': '' if first else f'=IF(OR(A{r}="",B{r}="",B{r - 1}=""),"",(B{r}-N(C{r}))/B{r - 1}-1)',
             'E': f'=IF(B{r}="","",1)' if first else f'=IF(D{r}="","",E{r - 1}*(1+D{r}))',
             'F': f'=IF(E{r}="","",MAX($E${NAV0}:E{r}))',
             'G': f'=IF(E{r}="","",1-E{r}/F{r})'}
        for col, v in f.items():
            if v:
                ws[f'{col}{r}'] = v
        for col, fmt in (('A', DATE), ('B', MONEY), ('C', MONEY)):
            c = ws[f'{col}{r}']
            c.fill, c.font, c.number_format = fill(INPUT_FILL), font(10, False, INPUT_FONT), fmt
        ws[f'D{r}'].number_format = '0.00%'
        ws[f'E{r}'].number_format = '0.0000'
        ws[f'F{r}'].number_format = '0.0000'
        ws[f'G{r}'].number_format = PCT
    ws.freeze_panes = 'A5'
    ws.sheet_properties.tabColor = NAVY
    name(wb, 'IBKR_NAVAnchor', f'Performance!$A${NAV0 - 1}')
    return ws


# ----------------------------------------------------------------- Dashboard
def build_dashboard(wb):
    ws = wb.create_sheet('Dashboard')
    title(ws, 'BAYESIAN CAPITAL  ·  25-SLOT MANAGED BOOK', '')
    put(ws, 'B2', '="Daily monitoring  ·  data to "&TEXT(MAX(Query!$A:$A),"ddd d mmm yyyy")&"  ·  today "'
        '&TEXT(TODAY(),"ddd d mmm yyyy")', f=font(10, False, '595959', True))
    lv1, sv1 = LV0 + NS - 1, SV0 + 2 * NS - 1
    tiles = [
        ('BOOK VALUE', '=E_Equity', MONEY, '="Peak "&TEXT(MAX(N(PF_PeakNAV),E_Equity),"$#,##0")'),
        ('DRAWDOWN FROM PEAK', '=IF(MAX(N(PF_PeakNAV),E_Equity)=0,0,1-E_Equity/MAX(N(PF_PeakNAV),E_Equity))',
         PCT, '="Worst since start "&TEXT(N(PF_MaxDD),"0.0%")'),
        ('CASH', '=P_Cash', MONEY, '=TEXT(IF(E_Equity=0,0,P_Cash/E_Equity),"0%")&" of the book"'),
        ('BREADTH GATE', '=IF(E_GateOn,"ON","off")', None,
         '=E_Below&" of "&E_CoreN&" Core below average · fires at "&P_GateK'),
        ('LENT TO BENCH', '=E_Lent', MONEY, '=COUNTIF(Orders!$G$%d:$G$%d,"BUY — Bench*")&" Bench orders"' % (OR0, OR0 + 2 * NS - 1)),
        ('ORDERS TO SEND', '=COUNTIF(Orders!$L$%d:$L$%d,"YES")' % (OR0, OR0 + 2 * NS - 1), INT,
         '=TEXT(E_Orders,"$#,##0")&" to send"'),
        ('LARGEST NAME', f'=MAX(Engine!$AK${LV0}:$AK${lv1})', PCT,
         f'=IF(MAX(Engine!$AK${LV0}:$AK${lv1})=0,"no positions",INDEX(Engine!$B${LV0}:$B${lv1},'
         f'MATCH(MAX(Engine!$AK${LV0}:$AK${lv1}),Engine!$AK${LV0}:$AK${lv1},0))&" · cap "&TEXT(P_ExpCap,"0%"))'),
        ('ALERTS', '=COUNTIF($C$10:$C$17,"▲*")', INT, '=IF(COUNTIF($C$10:$C$17,"▲*")=0,"all clear","see the list below")'),
    ]
    for i, (lab, f, fmt, sub) in enumerate(tiles):
        c1 = 2 + 2 * i
        rng = f'{CL(c1)}{{r}}:{CL(c1 + 1)}{{r}}'
        for rr, (val, ft, fm) in zip((4, 5, 6), ((lab, font(8, True, 'FFFFFF'), None),
                                                  (f, font(14, True, NAVY), fmt),
                                                  (sub, font(8, False, '595959'), None))):
            ws.merge_cells(rng.format(r=rr))
            c = ws[f'{CL(c1)}{rr}']
            c.value, c.font = val, ft
            if fm:
                c.number_format = fm
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=(rr == 6))
            for cc in (c1, c1 + 1):
                ws.cell(rr, cc).fill = fill(NAVY if rr == 4 else PALE)
                ws.cell(rr, cc).border = Border(left=Side(style='medium', color='FFFFFF') if cc == c1 else None)
        ws.cell(5, c1).fill = fill(PALE)
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 26
    ws.conditional_formatting.add('H5', CellIsRule(operator='equal', formula=['"ON"'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add('D5', CellIsRule(operator='greaterThan', formula=['0.1'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add('P5', CellIsRule(operator='greaterThan', formula=['0'],
                                  fill=fill(AMB_F), font=Font(color=AMB_T, bold=True)))
    # alerts
    bar(ws, 8, 2, 24, 'ALERTS')
    so, sh = f'Orders!$G${OR0}:$G${OR0 + 2 * NS - 1}', f'Engine!$AQ${LV0}:$AQ${lv1}'
    nstop = f'COUNTIF({so},"STOP*")'
    near = (f'SUMPRODUCT((Engine!$U${SV0}:$U${sv1}<>"")*(Engine!$U${SV0}:$U${sv1}>=P_StopDays-10)'
            f'*(Engine!$U${SV0}:$U${sv1}<P_StopDays))')
    disc = f'COUNTIFS(Log!$B${LG0}:$B${LG1},"Disc",Log!$V${LG0}:$V${LG1},"PAST STOP")'
    over = f'COUNTIF(Engine!$AK${LV0}:$AK${lv1},">"&P_ExpCap)'
    stale = f'COUNTIFS(Engine!$C${LV0}:$C${lv1},"<>Shadow",{sh},"<>OK")'
    veto = f'COUNTIF({so},"PM VETO*")'
    alerts = [
        '=IF(E_GateOn,"▲ Breadth gate ON — "&E_Below&" of "&E_CoreN&" Core names below their average; gated names hold cash"'
        '&IF(P_Lend="YES"," and lend it to the Bench.","."),"✓ Breadth gate off — "&E_Below&" of "&E_CoreN&" Core names below their average (fires at "&P_GateK&").")',
        f'=IF({nstop}>0,"▲ "&{nstop}&" position(s) at the time stop — they sell at the open.","✓ No position at its time stop.")',
        f'=IF({near}>0,"▲ "&{near}&" position(s) within 10 days of the time stop.","✓ Nothing within 10 days of its time stop.")',
        f'=IF({disc}>0,"▲ "&{disc}&" discretionary ticket(s) past their stop — close them.","✓ No discretionary ticket past its stop.")',
        f'=IF({over}>0,"▲ "&{over}&" name(s) above the "&TEXT(P_ExpCap,"0%")&" exposure cap — no new buys there.","✓ Every name inside the "&TEXT(P_ExpCap,"0%")&" exposure cap.")',
        f'=IF({stale}>0,"▲ "&{stale}&" live slot(s) with stale or missing data — refresh Query before trading.","✓ Every live slot has current data.")',
        f'=IF({veto}>0,"▲ "&{veto}&" order(s) vetoed by the pre-market rule.","✓ No pre-market vetoes.")',
        '=IF(P_NewBuys<>"YES","▲ NEW BUYS PAUSED (Parameters) — targets and stops still run.","✓ New buys allowed.")',
    ]
    for i, f in enumerate(alerts):
        put(ws, f'C{10 + i}', f, f=font(10))
    ws.conditional_formatting.add('C10:C17', FormulaRule(formula=['LEFT(C10,1)="▲"'],
                                  font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add('C10:C17', FormulaRule(formula=['LEFT(C10,1)="✓"'],
                                  font=Font(color='548235')))
    # slot table
    bar(ws, 19, 2, 24, 'SLOTS')
    lab = ['Slot', 'Ticker', 'Tier', 'Last close', 'Data to', 'vs 200-day', 'Gated', 'Bayes buy',
           'OU buy', "Today's orders $", 'Sleeves held', 'Oldest (days)', 'Stop in (days)',
           'Exposure', 'Liquidity cap $', 'Cap binding', 'Live fills (28 days)', 'Model fills (20 sessions)', 'Data',
           'Action (Bayes)', 'Action (OU)', 'Note']
    header(ws, 20, 2, lab, 30)
    for k in range(1, NS + 1):
        r = 20 + k
        lr = LV0 + k - 1
        s0, s1 = SV0 + 2 * (k - 1), SV0 + 2 * (k - 1) + 1
        o0, o1 = OR0 + 2 * (k - 1), OR0 + 2 * (k - 1) + 1
        f = {'B': k, 'C': f'=Engine!$B${lr}', 'D': f'=Engine!$C${lr}', 'E': f'=Engine!$H${lr}',
             'F': f'=Engine!$G${lr}', 'G': f'=Engine!$AB${lr}', 'H': f'=IF(Engine!$AD${lr},"GATED","")',
             'I': f'=Engine!$O${lr}', 'J': f'=Engine!$Y${lr}', 'K': f'=Engine!$Q${s0}+Engine!$Q${s1}',
             'L': f'=N(Engine!$H${s0})+N(Engine!$H${s1})',
             'M': f'=IF(L{r}=0,"",MAX(N(Engine!$U${s0}),N(Engine!$U${s1})))',
             'N': f'=IF(M{r}="","",P_StopDays-M{r})',
             'O': f'=Engine!$AK${lr}', 'P': f'=Engine!$AI${lr}', 'Q': f'=Engine!$AP${lr}',
             'R': f'=Engine!$AR${lr}', 'S': f'=Engine!$AS${lr}', 'T': f'=Engine!$AQ${lr}',
             'U': f'=Orders!$G${o0}', 'V': f'=Orders!$G${o1}', 'W': f'=""&Slots!$N${SL0 + k - 1}'}
        fm = {'E': PRICE, 'F': DATE, 'G': '+0.0%;-0.0%', 'I': PRICE, 'J': PRICE, 'K': MONEY, 'O': PCT,
              'P': MONEY}
        for col, v in f.items():
            put(ws, f'{col}{r}', v, fmt=fm.get(col), al=CENTER if col in 'BCDHLMNQRST' else None,
                f=font(9 if col == 'W' else 10, False, '595959' if col == 'W' else '000000'))
        for c in range(2, 24):
            ws.cell(r, c).border = Border(bottom=HAIR)
    t0, t1 = 21, 20 + NS
    tier_format(ws, f'D{t0}:D{t1}')
    ws.conditional_formatting.add(f'G{t0}:G{t1}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=RED_T)))
    ws.conditional_formatting.add(f'H{t0}:H{t1}', CellIsRule(operator='equal', formula=['"GATED"'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add(f'N{t0}:N{t1}', FormulaRule(formula=[f'AND(N{t0}<>"",N{t0}<=0)'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add(f'N{t0}:N{t1}', FormulaRule(formula=[f'AND(N{t0}<>"",N{t0}<=10)'],
                                  fill=fill(AMB_F), font=Font(color=AMB_T)))
    ws.conditional_formatting.add(f'O{t0}:O{t1}', FormulaRule(formula=[f'O{t0}>P_ExpCap'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    ws.conditional_formatting.add(f'O{t0}:O{t1}', FormulaRule(formula=[f'O{t0}>0.8*P_ExpCap'],
                                  fill=fill(AMB_F), font=Font(color=AMB_T)))
    ws.conditional_formatting.add(f'O{t0}:O{t1}', DataBarRule(start_type='num', start_value=0, end_type='num',
                                  end_value=0.3, color='5B9BD5'))
    ws.conditional_formatting.add(f'T{t0}:T{t1}', CellIsRule(operator='notEqual', formula=['"OK"'],
                                  fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
    for col in ('U', 'V'):
        rg = f'{col}{t0}:{col}{t1}'
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({col}{t0},3)="BUY"'],
                                      font=Font(color=GRN_T, bold=True)))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({col}{t0},4)="STOP"'],
                                      fill=fill(RED_F), font=Font(color=RED_T, bold=True)))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({col}{t0},6)="SHADOW"'],
                                      font=Font(color='808080', italic=True)))
    widths(ws, {'B': 6, 'C': 9, 'D': 8, 'E': 10, 'F': 10, 'G': 9, 'H': 8, 'I': 10, 'J': 10, 'K': 13,
                'L': 8, 'M': 8, 'N': 8, 'O': 10, 'P': 12, 'Q': 9, 'R': 8, 'S': 8, 'T': 8, 'U': 24,
                'V': 24, 'W': 60})
    ws.freeze_panes = 'E21'
    ws.sheet_properties.tabColor = NAVY
    return ws


# ----------------------------------------------------------------- Shadow
def build_shadow(wb):
    ws = wb.create_sheet('Shadow')
    title(ws, 'SHADOW PANEL — admission', 'Names tracked on paper, why each is not in the book, and the '
          'evidence against the five admission tests. Updated daily; tests 2–4 after each quarterly re-run.')
    put(ws, 'B3', 'Tests: 1 Role (β < max for the Bench; AI-related names only replace Core) · 2 Fit (settings '
        'carry over) · 3 Book test, both regimes · 4 Liquidity at current capital · 5 Shadow conduct '
        '(minimum days and paper fills, clean data).', f=font(9, False, '595959', True))
    lab = ['#', 'Ticker', 'Slot', 'Shadow since', 'Days', 'Why it is not in the book', 'β',
           'T1 Role', 'Fit pairs', 'T2 Fit', 'T3 Book test', 'T3 evidence', 'T4 Liquidity', 'T4 evidence',
           'Paper fills', 'Paper return', 'Stop-outs', 'Data', 'T5 Conduct', 'Passed', 'Status']
    header(ws, 5, 2, lab, 30)
    lv1 = LV0 + NS - 1
    for i in range(12):
        r = 6 + i
        n = i + 1
        sl = f'$D{r}'
        sr = f'(SL0-1+{sl})'
        idx = lambda col: f'=IF({sl}="","",INDEX(Slots!${col}${SL0}:${col}${SL0 + NS - 1},{sl}))'
        lidx = lambda col: f'=IF({sl}="","",INDEX(Engine!${col}${LV0}:${col}${lv1},{sl}))'
        f = {
            'B': f'=IF(D{r}="","",{n})',
            'D': f'=IFERROR(MATCH({n},Slots!$AF${SL0}:$AF${SL0 + NS - 1},0),"")',
            'C': idx('C'), 'E': idx('M'),
            'F': f'=IF(OR(D{r}="",E{r}=""),"",TODAY()-E{r})',
            'G': idx('N'), 'H': idx('G'),
            'I': f'=IF(D{r}="","",IF(H{r}<P_BetaMax,"PASS","CORE ONLY"))',
            'J': idx('J'), 'K': idx('O'), 'L': idx('P'), 'M': idx('Q'), 'N': idx('R'), 'O': idx('S'),
            'P': lidx('AU'), 'Q': lidx('AV'), 'R': lidx('AW'), 'S': lidx('AQ'),
            'T': f'=IF(D{r}="","",IF(AND(N(F{r})>=P_ShadowDays,N(P{r})>=P_ShadowFills,S{r}="OK"),"PASS","PENDING"))',
            'U': f'=IF(D{r}="","",(I{r}="PASS")+(K{r}="PASS")+(L{r}="PASS")+(N{r}="PASS")+(T{r}="PASS"))',
            'V': f'=IF(D{r}="","",IF(U{r}=5,"READY FOR REVIEW",U{r}&" of 5 passed"))',
        }
        fm = {'E': DATE, 'H': '0.00', 'Q': '+0.0%;-0.0%'}
        for col, v in f.items():
            put(ws, f'{col}{r}', v, fmt=fm.get(col),
                al=WRAP if col in 'GMO' else CENTER, f=font(9 if col in 'GMO' else 10))
        ws.row_dimensions[r].height = 44
        for c in range(2, 23):
            ws.cell(r, c).border = Border(bottom=HAIR)
    rng = 'I6:T17'
    for word, (bg, fg) in (('PASS', (GRN_F, GRN_T)), ('FAIL', (RED_F, RED_T)), ('CORE ONLY', (AMB_F, AMB_T)),
                           ('PENDING', (AMB_F, AMB_T)), ('NOT RUN', (GREY, '595959')),
                           ('NOT MEASURED', (GREY, '595959'))):
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{word}"'],
                                      fill=fill(bg), font=Font(color=fg, bold=(word == 'PASS'))))
    ws.conditional_formatting.add('V6:V17', CellIsRule(operator='equal', formula=['"READY FOR REVIEW"'],
                                  fill=fill(GRN_F), font=Font(color=GRN_T, bold=True)))
    widths(ws, {'B': 4, 'C': 8, 'D': 5, 'E': 10, 'F': 6, 'G': 52, 'H': 6, 'I': 10, 'J': 6, 'K': 9,
                'L': 10, 'M': 22, 'N': 11, 'O': 18, 'P': 7, 'Q': 8, 'R': 7, 'S': 8, 'T': 10, 'U': 7, 'V': 17})
    bar(ws, 19, 2, 22, 'WHAT ADMISSION NEEDS')
    rules = [
        'Shadow time checks execution and data, not edge: three months is far too short to judge a name on P&L (§3.99).',
        'Admission rests on the book test re-run each quarter on the extended data: adding the name must not cost more than ~1pp of test-half return or worsen the 2022 drawdown.',
        'A name that fails the fit test (T2) can still be admitted on neutral settings if tests 1, 3, 4 and 5 pass.',
        'An AI-related name (T1 = CORE ONLY) can enter only by replacing a Core name, after a book test of the new Core set.',
        'READY FOR REVIEW means the evidence is complete — the decision is still yours. Follow Guide procedure A and record it in the Change Log.',
    ]
    for i, t in enumerate(rules):
        put(ws, f'B{20 + i}', f'{i + 1}.  {t}', f=font(9, False, '404040'))
    ws.freeze_panes = 'D6'
    ws.sheet_properties.tabColor = NAVY
    return ws


# ----------------------------------------------------------------- Change Log
def build_changelog(wb):
    ws = wb.create_sheet('Change Log')
    title(ws, 'CHANGE LOG', 'Every decision that changes the book: tiers, tickers, parameters, admissions, '
          'removals. Dated, with the reason and the evidence — the record of active management.')
    header(ws, 5, 2, ['Date', 'Type', 'Slot', 'Ticker', 'From', 'To', 'Reason', 'Evidence', 'By'], 22)
    widths(ws, {'B': 11, 'C': 16, 'D': 6, 'E': 8, 'F': 14, 'G': 14, 'H': 60, 'I': 30, 'J': 8})
    seeds = [
        (dt.date(2026, 10, 23), 'Setup', '', '', '', '25 slots',
         'Book created: 12 Core, 5 Bench (NEM LEN DLTR RKT TEAM), 8 Shadow. Breadth gate K≥6 on the 200-day, '
         'held cash lent to the Bench, 25% exposure cap, liquidity cap (CF 2%, others 10%). Fresh trade log.',
         'HANDOVER §3.93–3.101', 'DF'),
        (dt.date(2026, 10, 23), 'Parameter', 8, 'CF', '—', '2%', 'Liquidity cap: CF is the tightest name at $10m.',
         '§3.100', 'DF'),
    ]
    for i, row in enumerate(seeds):
        r = 6 + i
        for j, v in enumerate(row):
            c = ws.cell(r, 2 + j, v)
            inp(c, DATE if j == 0 else None)
            c.alignment = WRAP
        ws.row_dimensions[r].height = 44 if i == 0 else 18
    for r in range(6 + len(seeds), 506):
        for j in range(9):
            inp(ws.cell(r, 2 + j), DATE if j == 0 else None)
    dv = DataValidation(type='list', formula1='"Setup,Tier change,Ticker change,Parameter,Admission,Removal,'
                        'Discretionary,Data,Other"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('C6:C505')
    ws.freeze_panes = 'B6'
    ws.sheet_properties.tabColor = '548235'
    return ws


# ----------------------------------------------------------------- Guide
GUIDE = [
    ('h', 'WHAT THIS BOOK DOES'),
    ('p', 'Twenty-five fixed slots. Core names trade every day as the live book always has. Bench names trade only when '
          'the breadth gate is on, with the cash it holds back from gated Core names. Shadow names run their model on '
          'paper, building the record used to decide admission. Off holds a slot\'s cash. The ticker and tier of a slot '
          'are the only things that change the book; the layout never moves.'),
    ('h', 'THE SHEETS'),
    ('t', [('Sheet', 'Purpose', 'You edit'),
           ('Dashboard', 'Daily monitoring: book value, drawdown, gate, lending, exposure, alerts, every slot at a glance.', 'nothing'),
           ('Orders', "Today's orders, one row per sleeve. Script 1 sends rows with Send = YES.", '9:00 price (or Script), overrides'),
           ('Log', 'Every fill, model and discretionary. Open positions drive holdings, exposure and stops.', 'fills (or Script 2)'),
           ('Shadow', 'Shadow names, why each is not in the book, and progress against the five admission tests.', 'nothing (inputs on Slots)'),
           ('Slots', 'The register: ticker, tier, settings and admission evidence per slot.', 'per the procedures below'),
           ('Parameters', 'Every rule the book runs on, with its value, evidence and review cadence.', 'on cadence, logged'),
           ('Change Log', 'Dated record of every decision.', 'every change'),
           ('Performance', 'Daily NAV log, time-weighted return, drawdown, P&L by slot and month.', 'NAV row (or Script 2)'),
           ('Engine', 'Levels, the gate and order sizing. Calculated.', 'nothing'),
           ('Query', 'Daily prices and volume by slot. Written by Script 1.', 'nothing'),
           ('M01–M25', 'Each slot\'s model (Kalman + OU), formula for formula as the live book, with a paper record.', 'nothing')]),
    ('h', 'COLOUR KEY'),
    ('k', [(INPUT_FILL, INPUT_FONT, 'Blue on cream: an input you (or a script) type.'),
           (SCRIPT_FILL, '375623', 'Green: written by Script 1 / Script 2, or linked from Parameters.'),
           (TIER['Core'][0], TIER['Core'][1], 'Core'), (TIER['Bench'][0], TIER['Bench'][1], 'Bench'),
           (TIER['Shadow'][0], TIER['Shadow'][1], 'Shadow'), (TIER['Off'][0], TIER['Off'][1], 'Off'),
           (RED_F, RED_T, 'Red: act today (stop due, over cap, stale data, gate on).'),
           (AMB_F, AMB_T, 'Amber: watch (within 10 days of the stop, near the cap, gated, capped, vetoed).'),
           (GRN_F, GRN_T, 'Green: an order to send, or a test passed.')]),
    ('h', 'DAILY ROUTINE'),
    ('n', ['Before 09:00 — Script 1 refreshes Query and writes Cash (and Net liquidation) to Parameters.',
           'Open the Dashboard. Clear every alert: stale data first (no trading on stale data), then stops due, then caps.',
           '09:00 — pre-market prices go into Orders (Script 1 or typed). Vetoed orders drop out and their cash is re-split.',
           'Review Orders. Use an override only for a reason you would record. Script 1 sends rows with Send = YES.',
           'During the day Script 2 writes fills to the Log with their target. Discretionary tickets go in as Type = Disc.',
           'After the close add the NAV row in Performance. Glance at the Shadow panel.']),
    ('h', 'REVIEW CADENCE'),
    ('n', ['Daily: tiers (only via the procedures), discretionary tickets, the emergency switch.',
           'Weekly: exposure and liquidity caps against the Dashboard; any name near the cap or with a binding liquidity cap.',
           'Monthly: Shadow panel — any READY FOR REVIEW; costs, interest rate and capital.',
           'Quarterly: re-run the admission tests (fit test over four window pairs, book test in both regimes, liquidity at '
           'current capital) and update Slots columns J and O–S; review the gate and order-rule parameters.',
           'Annually: the time stop and the model settings, only with walk-forward evidence.']),
    ('h', 'A.  PROMOTE A SHADOW NAME TO THE BENCH'),
    ('n', ['Shadow panel shows READY FOR REVIEW (all five tests PASS) — or you accept a named exception, in writing.',
           'Check T1 says PASS (β below the Bench maximum). An AI-related name cannot go to the Bench — see procedure B.',
           'Decide the settings basis: Fitted only if T2 = PASS; otherwise keep Neutral. Update Slots columns H and T–AC if fitted.',
           'Record the decision in the Change Log (Type = Admission, From Shadow, To Bench, reason, evidence).',
           'Change the slot\'s Tier to Bench on Slots. Nothing else moves.',
           'Next morning check Orders: the slot shows "BENCH — idle (gate off)" until the gate turns on.']),
    ('h', 'B.  PROMOTE A BENCH OR SHADOW NAME TO THE CORE (or replace a Core name)'),
    ('n', ['Run the book test of the NEW Core set (both regimes) and record the result on Slots.',
           'If the Core count changes, revisit the gate threshold: keep P_GateK at about half the Core count.',
           'Record in the Change Log (Type = Tier change) for every slot that moves.',
           'Set the outgoing Core name to Off first if it holds positions; let them run to target or stop.',
           'Change the incoming slot\'s Tier to Core. Next morning it bids with the rest of the Core.']),
    ('h', 'C.  DEMOTE OR SWITCH A SLOT OFF'),
    ('n', ['Set the Tier to Off. Its share of the morning cash is held in cash, never handed to other names (§3.98).',
           'Open positions are untouched: targets and the time stop still run.',
           'Use Off for genuine events — a halt, a corporate action, a removal decision — not for a fill forecast.',
           'Record it in the Change Log. When the slot is flat it can go to Shadow or take a new ticker (procedure D).']),
    ('h', 'D.  PUT A NEW TICKER IN A SLOT'),
    ('n', ['Only a slot that is Off or Shadow and holds no open positions.',
           'Type the new ticker on Slots, set Tier = Shadow, settings basis Neutral (Slots T–AC = the neutral vector), '
           'In Shadow since = today, and the research inputs you have (β, fit pairs, why it is a candidate).',
           'Have Script 1 back-fill the slot\'s Query block from 1 April 2024 (prices, dollar volume, first-hour volume).',
           'Check the slot\'s Model sheet shows data to the latest date, then record the change (Type = Ticker change).']),
    ('h', 'E.  A DISCRETIONARY TRADE'),
    ('n', ['Check the Dashboard exposure for the name: the ticket must leave it within the exposure cap.',
           'Log it with Type = Disc and Sleeve = DISC; target defaults to +P_DiscTarget.',
           'It is flagged PAST STOP after P_DiscStop sessions and OVER CAP if the name exceeds the cap. Close it on the flag.']),
    ('h', 'F.  CHANGE A PARAMETER'),
    ('n', ['Only on its review cadence, and only with the evidence the Parameters sheet names (gate and order rules '
           'need a backtest in both regimes).',
           'Change the blue value, update Last changed and Note, and add a Change Log row (Type = Parameter).']),
    ('h', 'G.  EMERGENCY STOP'),
    ('n', ['Set P_NewBuys = NO. Every new buy stops; targets and stops keep running; cash builds.',
           'Log why. Set it back to YES when the reason has cleared.']),
]


def build_guide(wb):
    ws = wb.create_sheet('Guide', 0)
    title(ws, 'GUIDE — how this book works and how to change it',
          'Read once in full; after that, the procedures A–G are the reference whenever a slot changes.')
    widths(ws, {'B': 4, 'C': 26, 'D': 92, 'E': 30})
    r = 4
    for kind, body in GUIDE:
        if kind == 'h':
            r += 1
            bar(ws, r, 2, 5, body)
            r += 1
        elif kind == 'p':
            ws.merge_cells(f'B{r}:E{r}')
            put(ws, f'B{r}', body, al=WRAPT)
            ws.row_dimensions[r].height = 58
            r += 1
        elif kind == 't':
            for i, row in enumerate(body):
                for j, v in enumerate(row):
                    c = put(ws, f'{CL(3 + j)}{r}', v, f=font(10, i == 0, NAVY if i == 0 else '000000'),
                            fl=LIGHT if i == 0 else None, al=WRAP)
                r += 1
        elif kind == 'k':
            for bg, fg, txt in body:
                put(ws, f'C{r}', 'Sample', f=font(10, True, fg), fl=bg, al=CENTER)
                put(ws, f'D{r}', txt)
                r += 1
        elif kind == 'n':
            for i, t in enumerate(body, start=1):
                put(ws, f'B{r}', i, f=font(10, True, BLUE), al=Alignment(horizontal='right', vertical='top'))
                ws.merge_cells(f'C{r}:E{r}')
                put(ws, f'C{r}', t, al=WRAP)
                ws.row_dimensions[r].height = 30 if len(t) > 110 else 16
                r += 1
    ws.sheet_properties.tabColor = NAVY
    return ws


# ----------------------------------------------------------------- main
def main():
    live, seed_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
    seed = json.load(open(seed_path))
    assert len(seed['slots']) == NS
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_dashboard(wb)
    build_orders(wb)
    build_log(wb)
    build_shadow(wb)
    build_slots(wb, seed)
    build_parameters(wb)
    build_changelog(wb)
    build_performance(wb)
    build_engine(wb)
    _, nd = build_query(wb, seed)
    assert nd <= N
    tpl, rows = model_templates(live)
    for k in range(1, NS + 1):
        build_model(wb, k, tpl, rows)
        print(f'  model sheet {msheet(k)} written', flush=True)
    build_guide(wb)
    unbalanced = []
    for ws in wb.worksheets:
        if ws.title.startswith('M') and ws.title != 'M01':
            continue                       # M02..M25 share M01's formulas
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    depth, instr = 0, False
                    for ch in v:
                        if ch == '"':
                            instr = not instr
                        elif not instr:
                            depth += (ch == '(') - (ch == ')')
                            if depth < 0:
                                break
                    if depth != 0 or instr:
                        unbalanced.append(f'{ws.title}!{c.coordinate}')
    assert not unbalanced, f'unbalanced formulas: {unbalanced[:10]}'
    from openpyxl.worksheet.properties import PageSetupProperties
    for ws in wb.worksheets:
        if ws.title in ('Guide', 'Dashboard', 'Orders', 'Shadow', 'Slots', 'Parameters', 'Change Log',
                        'Performance', 'Log'):
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_options.horizontalCentered = True
            ws.page_margins.left = ws.page_margins.right = 0.4
            ws.page_margins.top = ws.page_margins.bottom = 0.5
            ws.oddFooter.center.text = '&A  ·  page &P of &N'
            ws.oddFooter.center.size = 8
    wb['Orders'].print_title_rows = f'{OR0 - 1}:{OR0 - 1}'
    wb['Log'].print_title_rows = f'{LG0 - 1}:{LG0 - 1}'
    wb['Dashboard'].print_title_rows = '20:20'
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.active = 1
    wb.save(out)
    print(f'wrote {out}: {len(wb.sheetnames)} sheets, {nd} data rows of {N} capacity')


if __name__ == '__main__':
    main()
