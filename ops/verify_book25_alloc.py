"""
verify_book25_alloc.py — stress the 25-slot workbook's sizing rules against an
independent Python implementation of the same rules (HANDOVER 3.102).

Scenario written into a copy of the built workbook, then recalculated by
LibreOffice:
  cash $2.4m · gate threshold K = 3 (so the gate fires on current data) ·
  exposure cap 10% · CF switched Off · open positions: TSM Bayes (model),
  MU OU (model, bought 1 Sep: past its 50-day stop), VRT (discretionary) ·
  a 09:00 price on GM Bayes that vetoes it.

The Python side recomputes, from the recalculated sheet's levels and the Log,
every sleeve's category, base order, exposure/liquidity caps, final dollars
and shares, plus the gate count from Query closes, and compares.

Usage (from premarket_study/):
    python3 ../ops/verify_book25_alloc.py <built.xlsx> <recalc_dir> <seed_recalc.xlsx>
"""
import datetime as dt
import math
import os
import shutil
import subprocess
import sys

import openpyxl

NS, SL0, LV0, SV0, OR0, LG0 = 25, 7, 22, 51, 9, 7
CASH, K, CAP, FEE, THR = 2_400_000, 3, 0.10, 0.005, 0.04
TODAY = dt.date.today()


def pcell(wb, nm):
    ref = wb.defined_names[nm].attr_text
    sh, c = ref.split('!')
    return wb[sh][c.replace('$', '')]


def main():
    built, rdir, prev = sys.argv[1:4]
    pv = openpyxl.load_workbook(prev, data_only=True)
    gm_level = pv['Engine'].cell(LV0 + 5, 15).value          # GM Bayes level (slot 6, col O)
    scen = os.path.join(rdir, 'scenario.xlsx')
    wb = openpyxl.load_workbook(built)
    pcell(wb, 'P_Cash').value = CASH
    pcell(wb, 'P_GateK').value = K
    pcell(wb, 'P_ExpCap').value = CAP
    wb['Slots'][f'D{SL0 + 7}'] = 'Off'                       # CF
    log = wb['Log']
    rows = [('Model', 'TSM', 'Bayes', dt.date(2026, 10, 1), 400.0, 1000, 410.0),
            ('Model', 'MU', 'OU', dt.date(2026, 9, 1), 900.0, 500, 930.0),
            ('Disc', 'VRT', 'DISC', dt.date(2026, 10, 13), 220.0, 3000, 231.0)]
    for i, (ty, tk, sl, d, px, sh, tg) in enumerate(rows):
        r = LG0 + i
        for col, v in zip('BCDEFGI', (ty, tk, sl, d, px, sh, tg)):
            log[f'{col}{r}'] = v
    pm_gm = round(gm_level * 1.10, 2)
    wb['Orders'][f'M{OR0 + 2 * 5}'] = pm_gm                  # GM Bayes 09:00 price
    wb.save(scen)
    out = os.path.join(rdir, 'out')
    os.makedirs(out, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
                    '--outdir', out, scen], check=True, capture_output=True, timeout=1800)
    wv = openpyxl.load_workbook(os.path.join(out, 'scenario.xlsx'), data_only=True)
    en, q, o = wv['Engine'], wv['Query'], wv['Orders']

    # ---- independent recomputation
    lv = {}
    for k in range(1, NS + 1):
        r = LV0 + k - 1
        g = lambda c: en.cell(r, c).value
        closes = [q.cell(rr, 2 + 6 * (k - 1) + 3).value for rr in range(2, q.max_row + 1)]
        closes = [c for c in closes if isinstance(c, (int, float))]
        below = len(closes) >= 200 and closes[-1] < sum(closes[-200:]) / 200
        lv[k] = dict(tk=g(2), tier=g(3), close=g(8), bayes=g(15), ou=g(25), pi_b=g(13), pi_o=g(23),
                     below=below, fh=g(33), L=g(34))
    open_ = [(r[0], r[1], r[2], r[5]) for r in rows]          # type, ticker, sleeve, shares
    close_of = {v['tk']: v['close'] for v in lv.values()}
    equity = CASH + sum(sh * close_of[tk] for _, tk, _, sh in open_)
    nbelow = sum(1 for v in lv.values() if v['tier'] == 'Core' and v['below'])
    gate = nbelow >= K
    held_model = {(tk, sl) for ty, tk, sl, _ in open_ if ty == 'Model'}
    held_val = {}
    for _, tk, _, sh in open_:
        held_val[tk] = held_val.get(tk, 0) + sh * close_of[tk]
    cat = {}
    for k, v in lv.items():
        for sl, lvl in (('Bayes', v['bayes']), ('OU', v['ou'])):
            pm = pm_gm if (k == 6 and sl == 'Bayes') else None
            if (v['tk'], sl) in held_model:
                c = 'held'
            elif not (isinstance(lvl, (int, float)) and lvl > 0):
                c = 'no bid'
            elif v['tier'] == 'Shadow':
                c = 'shadow'
            elif v['tier'] == 'Off':
                c = 'off'
            elif pm is not None and lvl < pm * (1 - THR):
                c = 'veto'
            elif v['tier'] == 'Bench':
                c = 'bench'
            elif gate and v['below']:
                c = 'gated'
            else:
                c = 'core'
            cat[(k, sl)] = (c, lvl)
    n = {c: sum(1 for x in cat.values() if x[0] == c) for c in ('core', 'gated', 'off', 'bench')}
    den = n['core'] + n['gated'] + n['off']
    share = CASH / den if den else 0
    lend = CASH * n['gated'] / den if den else 0
    bench_each = min(lend / n['bench'], share) if n['bench'] else 0
    base = {key: (share if c == 'core' else bench_each if c == 'bench' else 0) for key, (c, _) in cat.items()}
    final, shares = {}, {}
    for k, v in lv.items():
        tot = base[(k, 'Bayes')] + base[(k, 'OU')]
        room = max(0, CAP * equity - held_val.get(v['tk'], 0))
        cap = room if v['fh'] in (None, '') else min(room, v['L'] * v['fh'])
        sc = 1 if tot <= 0 else min(1, cap / tot)
        for sl in ('Bayes', 'OU'):
            f_ = base[(k, sl)] * sc
            lvl = cat[(k, sl)][1]
            final[(k, sl)] = f_
            shares[(k, sl)] = math.floor(f_ / (lvl + FEE)) if f_ > 0 else 0

    # ---- compare
    print(f'scenario: equity ${equity:,.0f} · Core below {nbelow} (workbook {en["C6"].value}) · '
          f'gate {gate} (workbook {en["C7"].value}) · den {den} · share ${share:,.0f} · lend ${lend:,.0f} · '
          f'bench each ${bench_each:,.0f}')
    bad = 0
    assert abs(en['C4'].value - equity) < 0.01, (en['C4'].value, equity)
    assert en['C6'].value == nbelow and bool(en['C7'].value) == gate
    for k in range(1, NS + 1):
        for j, sl in enumerate(('Bayes', 'OU')):
            r = SV0 + 2 * (k - 1) + j
            wc, wf, ws_ = en.cell(r, 12).value, en.cell(r, 15).value, en.cell(r, 16).value
            pc = cat[(k, sl)][0]
            okk = wc == pc and abs((wf or 0) - final[(k, sl)]) < 0.01 and (ws_ or 0) == shares[(k, sl)]
            bad += not okk
            act = o.cell(OR0 + 2 * (k - 1) + j, 7).value
            if not okk or pc not in ('core', 'shadow'):
                print(f'  {"OK " if okk else "BAD"} {lv[k]["tk"]:5s} {sl:5s} {pc:7s} wb={wc:7s} '
                      f'${final[(k, sl)]:>10,.0f} vs ${wf or 0:>10,.0f}  {shares[(k, sl)]:>6d} vs {ws_ or 0:>6}  · {act}')
    sent = sum(en.cell(SV0 + i, 17).value or 0 for i in range(2 * NS))
    held_back = CASH - sent
    print(f'  orders ${sent:,.0f} + held back ${held_back:,.0f} = cash ${CASH:,.0f}; never above cash: {sent <= CASH}')
    print('ALLOCATION CHECK PASS' if bad == 0 and sent <= CASH else f'ALLOCATION CHECK FAILED ({bad})')


if __name__ == '__main__':
    main()
