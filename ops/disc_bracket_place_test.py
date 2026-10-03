"""
disc_bracket_place_test.py — exercise every part of the bracket placer except
the ib_insync calls, which need a gateway.

Builds a synthetic workbook with the live sheet's geometry and literal values
(the real one cannot be used: openpyxl blanks cached formula values on save,
so a freshly-patched workbook reads target=None until Excel has opened and
saved it -- which the placer correctly refuses on, and which is asserted here).

Run: python disc_bracket_place_test.py
"""
import datetime as dt
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import disc_bracket_place as P

TODAY = dt.date(2026, 10, 6)
SESSIONS = [dt.date(2026, 9, 21) + dt.timedelta(days=i) for i in range(30)
            if (dt.date(2026, 9, 21) + dt.timedelta(days=i)).weekday() < 5]


def build(rows):
    wb = openpyxl.Workbook()
    q = wb.active
    q.title = 'Query'
    names = ['TSM', 'VRT', 'VST']
    q.cell(row=1, column=1, value='Date')
    for i, nm in enumerate(names):
        for j, f in enumerate('OHLC'):
            q.cell(row=1, column=2 + 4 * i + j, value=f'{nm}_{f}')
    px = dict(TSM=470.0, VRT=250.0, VST=140.0)
    for r, d in enumerate(SESSIONS, start=2):
        q.cell(row=r, column=1, value=d)
        for i, nm in enumerate(names):
            for j, v in enumerate((px[nm], px[nm] * 1.01, px[nm] * 0.99, px[nm])):
                q.cell(row=r, column=2 + 4 * i + j, value=v)
    at = wb.create_sheet('Active Trading')
    for i, t in enumerate(rows):
        r = P.ROW0 + i
        for k, c in P.COL.items():
            if k in t:
                at.cell(row=r, column=c, value=t[k])
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def run(rows, journal=None, **kw):
    path = build(rows)
    try:
        tickets, sessions = P.read_tickets(path)
        return P.screen(tickets, journal or {}, kw.pop('when', TODAY),
                        kw.pop('max_notional', 500_000), kw.pop('allow_far', False))
    finally:
        os.unlink(path)


def case(name, rows, expect_placed, expect_reason=None, **kw):
    place, refuse = run(rows, **kw)
    ok = len(place) == expect_placed
    if expect_reason is not None:
        ok = ok and any(expect_reason in (w or '') for _, w in refuse)
    print(f'  {"PASS" if ok else "FAIL"}  {name}')
    if not ok:
        print(f'        placed={len(place)} refused={[(t["stock"], w) for t, w in refuse]}')
    return ok


GOOD = dict(stock='VRT', tranche='DISC', buyd=TODAY, buypx=240.0,
            shares=1000, target=252.0)


def main():
    print('bracket placer — screening logic')
    ok = []
    ok.append(case('a clean ticket is placed', [GOOD], 1))
    ok.append(case('a sold ticket is ignored',
                   [dict(GOOD, selld=TODAY)], 0))
    ok.append(case("another day's ticket is ignored",
                   [dict(GOOD, buyd=dt.date(2026, 10, 2))], 0))
    ok.append(case('a blank target is refused',
                   [{k: v for k, v in GOOD.items() if k != 'target'}], 0,
                   'no target price'))
    ok.append(case('a target below the limit is refused',
                   [dict(GOOD, target=230.0)], 0, 'not above the limit'))
    ok.append(case('the sheet check column vetoes',
                   [dict(GOOD, check='CHECK TICKER')], 0, 'sheet check'))
    ok.append(case('over the per-ticket notional cap',
                   [dict(GOOD, shares=100_000)], 0, 'over the per-ticket'))
    ok.append(case('a limit far from the last close is refused',
                   [dict(GOOD, buypx=24.0, target=25.2)], 0, 'from the last close'))
    ok.append(case('  ... and allowed with --allow-far',
                   [dict(GOOD, buypx=24.0, target=25.2)], 1, allow_far=True))

    # idempotency: the journal key must stop a second placement
    path = build([GOOD])
    try:
        tickets, _ = P.read_tickets(path)
        t = [x for x in tickets if x['buy_date'] == TODAY][0]
        j = {P.key(t): dict(placed_at='2026-10-06T09:31:00')}
        place, refuse = P.screen(tickets, j, TODAY, 500_000, False)
        good = len(place) == 0 and any('already placed' in (w or '') for _, w in refuse)
        print(f'  {"PASS" if good else "FAIL"}  a journalled ticket is not placed twice')
        ok.append(good)
    finally:
        os.unlink(path)

    # the openpyxl trap: formulas with no cached value must refuse, not guess
    wb = openpyxl.Workbook()
    wb.active.title = 'Query'
    wb['Query'].cell(row=1, column=1, value='Date')
    at = wb.create_sheet('Active Trading')
    for k, c in P.COL.items():
        if k in GOOD:
            at.cell(row=P.ROW0, column=c, value=GOOD[k])
    at.cell(row=P.ROW0, column=P.COL['target'], value='=X42*1.05')
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    try:
        tickets, _ = P.read_tickets(path)
        place, refuse = P.screen(tickets, {}, TODAY, 500_000, False)
        good = len(place) == 0 and any('no target price' in (w or '') for _, w in refuse)
        print(f'  {"PASS" if good else "FAIL"}  an uncalculated workbook refuses '
              f'rather than guessing')
        ok.append(good)
    finally:
        os.unlink(path)

    # time stop lands on the Nth session after entry
    ts = P.time_stop_date(SESSIONS, SESSIONS[0], 10)
    good = ts == SESSIONS[10]
    print(f'  {"PASS" if good else "FAIL"}  the time stop is the 10th session after entry')
    ok.append(good)

    print(f'\n{sum(ok)}/{len(ok)} passed')
    return 0 if all(ok) else 1


if __name__ == '__main__':
    sys.exit(main())
