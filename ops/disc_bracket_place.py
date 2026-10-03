"""
disc_bracket_place.py — place the discretionary log's tickets as IBKR brackets
(user, 3 Oct 2026).

The discretionary carve-out only works if a ticket is one action: the log
already derives the exit from the Tranche label (wire_disc_targets.py --
"Bayes"/"OU" take that stock's sleeve premium off the previous close, "DISC"
takes the standard 5% off the entry), and this places the matching bracket so
the exit is defined at entry and never revisited.

  parent   BUY  LMT  `Shares` of `Stock` at `Buy price`        (col X)
  child    SELL LMT  the same quantity at `Target sale`, GTC   (col AA)

transmitted together, so the take-profit is live the moment the buy fills.

IT NEVER WRITES THE WORKBOOK. openpyxl blanks every cached formula value on
save (HANDOVER section 4), which would leave the live sheet showing nothing
until Excel reopened it -- fatal for a workbook being traded from. Placement
state lives in a JSON journal beside this script instead, which is also what
makes re-running safe: a ticket already placed is skipped by (stock, date,
price, shares), so running twice does not double an order.

SAFETY RAILS, all on by default:
  * dry run unless --live is passed; the dry run prints the exact orders
  * --max-notional per ticket and --max-total per run, both enforced before
    anything is transmitted
  * tickets whose Check column says CHECK TICKER are refused (the entry price
    sits outside that ticker's own range that day -- six rows of the live log
    were VST trades typed as VRT, and with the exit now derived from the
    ticker that typo becomes a wrong order, not just a mislabelled one)
  * a ticket whose limit is far from the last close is refused unless
    --allow-far (catches a decimal slip)
  * the port is reported, and --live against a known LIVE port demands
    --i-mean-it

NOT AUTOMATED, deliberately: the ticket's 10-session time stop. IBKR has no
order type for "sell at the open after N sessions", so forcing it would mean a
scheduled market sell -- a second, riskier automation. `--due` reports which
open tickets have passed their time stop so it stays a human action.

UNTESTED AGAINST A BROKER. Everything except the ib_insync calls is exercised
by disc_bracket_place_test.py. Run it against IB PAPER first and reconcile the
fills by hand before pointing it at the live account.

Usage:
    python disc_bracket_place.py <workbook.xlsx>                 # dry run
    python disc_bracket_place.py <workbook.xlsx> --live          # place
    python disc_bracket_place.py <workbook.xlsx> --date 2026-10-06
    python disc_bracket_place.py <workbook.xlsx> --due           # time stops
"""
import argparse
import datetime as dt
import json
import os
import sys

import openpyxl

SHEET = 'Active Trading'
ROW0, ROW1 = 42, 231
COL = dict(stock=21, tranche=22, buyd=23, buypx=24, shares=25,
           target=27, selld=28, check=34)
JOURNAL = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       'disc_orders_journal.json')
PAPER_PORTS, LIVE_PORTS = {7497, 4002}, {7496, 4001}
TIME_STOP_SESSIONS = 10
FAR_FROM_CLOSE = 0.25          # a limit this far from the last close is a typo


def d_of(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def read_tickets(path):
    """Open discretionary tickets, with the last close for each stock."""
    wb = openpyxl.load_workbook(path, data_only=True)
    at = wb[SHEET]
    q = wb['Query']
    heads = {q.cell(row=1, column=c).value: c for c in range(1, q.max_column + 1)}
    closes, sessions = {}, []
    for r in range(2, q.max_row + 1):
        d = d_of(q.cell(row=r, column=1).value)
        if d is None:
            continue
        sessions.append(d)
        for nm in {h.split('_')[0] for h in heads if h and h.endswith('_C')}:
            v = q.cell(row=r, column=heads[f'{nm}_C']).value
            if isinstance(v, (int, float)):
                closes[nm] = (d, v)
    out = []
    for r in range(ROW0, ROW1 + 1):
        g = {k: at.cell(row=r, column=c).value for k, c in COL.items()}
        if not g['stock'] or g['buypx'] in (None, '') or not g['shares']:
            continue
        out.append(dict(
            row=r, stock=str(g['stock']).strip(),
            tranche=str(g['tranche'] or '').strip(),
            buy_date=d_of(g['buyd']), limit=float(g['buypx']),
            shares=int(g['shares']),
            target=float(g['target']) if isinstance(g['target'], (int, float)) else None,
            sold=d_of(g['selld']) is not None,
            check=str(g['check']).strip() if g['check'] else '',
            last_close=closes.get(str(g['stock']).strip(), (None, None))[1]))
    wb.close()
    return out, sorted(set(sessions))


def key(t):
    return f"{t['stock']}|{t['buy_date']}|{t['limit']:.4f}|{t['shares']}"


def load_journal():
    if os.path.exists(JOURNAL):
        with open(JOURNAL) as f:
            return json.load(f)
    return {}


def save_journal(j):
    with open(JOURNAL, 'w') as f:
        json.dump(j, f, indent=1, default=str)


def time_stop_date(sessions, buy_date, n=TIME_STOP_SESSIONS):
    later = [d for d in sessions if d > buy_date]
    return later[n - 1] if len(later) >= n else None


def screen(tickets, journal, when, max_notional, allow_far):
    """Split today's tickets into placeable and refused, with a reason each."""
    place, refuse = [], []
    for t in tickets:
        if t['sold']:
            continue
        if t['buy_date'] != when:
            continue
        why = None
        if key(t) in journal:
            why = f"already placed {journal[key(t)].get('placed_at', '')}"
        elif t['check']:
            why = f"sheet check says {t['check']!r}"
        elif t['target'] is None:
            why = 'no target price in the log'
        elif t['target'] <= t['limit']:
            why = f"target {t['target']:.2f} is not above the limit {t['limit']:.2f}"
        elif t['shares'] <= 0:
            why = 'non-positive share count'
        elif t['limit'] * t['shares'] > max_notional:
            why = (f'notional {t["limit"]*t["shares"]:,.0f} over the per-ticket '
                   f'cap {max_notional:,.0f}')
        elif (not allow_far and t['last_close']
                and abs(t['limit'] / t['last_close'] - 1) > FAR_FROM_CLOSE):
            why = (f"limit {t['limit']:.2f} is {abs(t['limit']/t['last_close']-1)*100:.0f}% "
                   f"from the last close {t['last_close']:.2f} — check for a typo")
        (refuse if why else place).append((t, why))
    return place, refuse


class IBBroker:
    """The only part that talks to IBKR. Isolated so the rest is testable and
    so swapping ib_insync for the native ibapi touches one class."""

    def __init__(self, host, port, client_id):
        from ib_insync import IB
        self.ib = IB()
        self.ib.connect(host, port, clientId=client_id)

    def place_bracket(self, stock, shares, limit, target):
        from ib_insync import Stock, LimitOrder
        c = Stock(stock, 'SMART', 'USD')
        self.ib.qualifyContracts(c)
        parent = LimitOrder('BUY', shares, round(limit, 2),
                            orderId=self.ib.client.getReqId(),
                            transmit=False, tif='DAY')
        child = LimitOrder('SELL', shares, round(target, 2),
                           orderId=self.ib.client.getReqId(),
                           parentId=parent.orderId, transmit=True, tif='GTC')
        trades = [self.ib.placeOrder(c, parent), self.ib.placeOrder(c, child)]
        self.ib.sleep(1)
        return [t.order.orderId for t in trades]

    def close(self):
        self.ib.disconnect()


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument('workbook')
    ap.add_argument('--live', action='store_true', help='actually transmit')
    ap.add_argument('--i-mean-it', action='store_true',
                    help='required with --live against a live-account port')
    ap.add_argument('--date', help='ticket buy date to place (default today)')
    ap.add_argument('--due', action='store_true',
                    help='report open tickets past their time stop and exit')
    ap.add_argument('--host', default='127.0.0.1')
    ap.add_argument('--port', type=int, default=7497, help='7497/4002 paper')
    ap.add_argument('--client-id', type=int, default=17)
    ap.add_argument('--max-notional', type=float, default=500_000)
    ap.add_argument('--max-total', type=float, default=1_500_000)
    ap.add_argument('--allow-far', action='store_true')
    a = ap.parse_args(argv)

    tickets, sessions = read_tickets(a.workbook)
    journal = load_journal()
    when = dt.date.fromisoformat(a.date) if a.date else dt.date.today()

    if a.due:
        print('OPEN TICKETS PAST THEIR TIME STOP '
              f'({TIME_STOP_SESSIONS} sessions; selling is yours, not the script\'s)')
        any_due = False
        for t in tickets:
            if t['sold'] or not t['buy_date']:
                continue
            ts = time_stop_date(sessions, t['buy_date'])
            if ts and ts <= sessions[-1]:
                any_due = True
                print(f"  row {t['row']:3d}  {t['stock']:5s} {t['tranche']:5s} "
                      f"bought {t['buy_date']} @ {t['limit']:.2f}  "
                      f"time stop was {ts}")
        if not any_due:
            print('  none.')
        return 0

    place, refuse = screen(tickets, journal, when, a.max_notional, a.allow_far)
    total = sum(t['limit'] * t['shares'] for t, _ in place)
    mode = 'LIVE' if a.live else 'DRY RUN'
    port_kind = ('PAPER' if a.port in PAPER_PORTS else
                 'LIVE ACCOUNT' if a.port in LIVE_PORTS else 'unrecognised')
    print(f'{mode} — tickets dated {when}, port {a.port} ({port_kind})\n')

    if refuse:
        print('REFUSED:')
        for t, why in refuse:
            print(f"  row {t['row']:3d}  {t['stock']:5s} {t['shares']:>6d} @ "
                  f"{t['limit']:.2f}   {why}")
        print()
    if not place:
        print('nothing to place.')
        return 0

    print('TO PLACE (bracket: BUY LMT, then SELL LMT GTC at the target):')
    for t, _ in place:
        print(f"  row {t['row']:3d}  {t['stock']:5s} {t['tranche']:5s} "
              f"BUY {t['shares']:>6d} @ {t['limit']:>9.2f}  ->  SELL @ "
              f"{t['target']:>9.2f}  ({(t['target']/t['limit']-1)*100:+.2f}%)  "
              f"notional {t['limit']*t['shares']:>12,.0f}")
    print(f'\n  total notional {total:,.0f} against a run cap of {a.max_total:,.0f}')
    if total > a.max_total:
        print('  REFUSED: over the run cap. Nothing placed.')
        return 1
    if not a.live:
        print('\ndry run — nothing was sent. Re-run with --live to transmit.')
        return 0
    if a.port in LIVE_PORTS and not a.i_mean_it:
        print('\nREFUSED: --live against a live-account port needs --i-mean-it.')
        return 1

    broker = IBBroker(a.host, a.port, a.client_id)
    try:
        for t, _ in place:
            ids = broker.place_bracket(t['stock'], t['shares'], t['limit'], t['target'])
            journal[key(t)] = dict(row=t['row'], stock=t['stock'],
                                   tranche=t['tranche'], buy_date=str(t['buy_date']),
                                   limit=t['limit'], shares=t['shares'],
                                   target=t['target'], order_ids=ids,
                                   time_stop=str(time_stop_date(sessions, t['buy_date'])),
                                   placed_at=dt.datetime.now().isoformat(timespec='seconds'))
            save_journal(journal)
            print(f"  placed {t['stock']:5s} order ids {ids}")
    finally:
        broker.close()
    print(f'\njournal: {JOURNAL}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
