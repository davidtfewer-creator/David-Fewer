"""
disc_reentry_live.py — what the live discretionary log actually did (3 Oct 2026).

The user's observation: when BOTH sleeves of a name are held, that name is
frozen out of the book until one sleeve sells, and a violently-moving name
can offer several round trips in that window. He has been taking those by
hand in the discretionary log and reports strong margin.

This reads the live workbook and asks, per discretionary trade, the only
question that matters before formalising anything:

  Was the name's model capacity ACTUALLY frozen when the trade was placed?

Sleeve occupancy is reconstructed from the main blotter's own buy/sell dates
(Active Trading A41:L..., one row per completed or open model trade), so the
answer comes from the book's own record rather than from recollection.

Reports: realised and open P&L per trade, hold in sessions, the occupancy
state of that name's two sleeves on the entry date, and the split between
re-entries into a frozen name and discretionary trades that needed no
freeze at all.

Usage: python disc_reentry_live.py <workbook.xlsx>
"""
import datetime as dt
import json
import statistics
import sys

import openpyxl

OUT = 'disc_reentry_live.json'
MAIN_HDR, DISC_HDR = 41, 41          # header rows of the two logs
DISC_COL0 = 20                       # the discretionary log starts at column T


def d_of(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def read_logs(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    at = wb['Active Trading']

    main = []
    for r in range(MAIN_HDR + 1, at.max_row + 1):
        nm = at.cell(row=r, column=2).value
        if not nm:
            continue
        main.append(dict(
            row=r, name=str(nm).strip(), sleeve=str(at.cell(row=r, column=3).value or '').strip(),
            buy=d_of(at.cell(row=r, column=4).value), px=at.cell(row=r, column=5).value,
            shares=at.cell(row=r, column=6).value, cost=at.cell(row=r, column=7).value,
            sell=d_of(at.cell(row=r, column=8).value), spx=at.cell(row=r, column=9).value,
            pnl=at.cell(row=r, column=11).value))

    disc = []
    c = DISC_COL0
    for r in range(DISC_HDR + 1, at.max_row + 1):
        nm = at.cell(row=r, column=c + 1).value
        if not nm:
            continue
        disc.append(dict(
            row=r, name=str(nm).strip(), label=str(at.cell(row=r, column=c + 2).value or '').strip(),
            buy=d_of(at.cell(row=r, column=c + 3).value), px=at.cell(row=r, column=c + 4).value,
            shares=at.cell(row=r, column=c + 5).value, cost=at.cell(row=r, column=c + 6).value,
            target=at.cell(row=r, column=c + 7).value, sell=d_of(at.cell(row=r, column=c + 8).value),
            spx=at.cell(row=r, column=c + 9).value, proceeds=at.cell(row=r, column=c + 10).value,
            pnl=at.cell(row=r, column=c + 11).value))

    # the current book state, for open positions and the live freeze picture
    orders = {}
    for r in range(19, 37):
        nm, tr = at.cell(row=r, column=1).value, at.cell(row=r, column=2).value
        if nm and tr:
            orders[(str(nm).strip(), str(tr).strip())] = dict(
                status=at.cell(row=r, column=3).value, note=at.cell(row=r, column=9).value,
                close=at.cell(row=r, column=10).value, target=at.cell(row=r, column=7).value)

    # the log is hand-typed and one session's rows carry the wrong ticker.
    # Reassign any trade whose buy price sits outside that name's own range on
    # the day to the only book name whose range contains it.
    relabelled = []
    q = wb['Query']
    heads = {q.cell(row=1, column=i).value: i for i in range(1, q.max_column + 1)}
    px = {}
    for nm in sorted({h.split('_')[0] for h in heads if h and '_' in h}):
        ser = {}
        for r in range(2, q.max_row + 1):
            d = d_of(q.cell(row=r, column=1).value)
            v = q.cell(row=r, column=heads[f'{nm}_C']).value
            if d and isinstance(v, (int, float)):
                ser[d] = v
        px[nm] = ser
    rng = {}
    for nm in px:
        ser = {}
        for r in range(2, q.max_row + 1):
            d = d_of(q.cell(row=r, column=1).value)
            lo = q.cell(row=r, column=heads[f'{nm}_L']).value
            hi = q.cell(row=r, column=heads[f'{nm}_H']).value
            if d and isinstance(lo, (int, float)):
                ser[d] = (lo, hi)
        rng[nm] = ser
    for t in disc:
        own = rng.get(t['name'], {}).get(t['buy'])
        if not own or not t['px'] or own[0] * 0.98 <= t['px'] <= own[1] * 1.02:
            continue
        cands = [nm for nm, s2 in rng.items()
                 if t['buy'] in s2 and s2[t['buy']][0] <= t['px'] <= s2[t['buy']][1]]
        if len(cands) == 1:
            relabelled.append((t['row'], t['name'], cands[0], t['buy'], t['px']))
            t['name'] = cands[0]
    if relabelled:
        print('TICKER CORRECTIONS (buy price outside the logged name\'s own range '
              'that day, inside exactly one other name\'s):')
        for row, was, now, d, p_ in relabelled:
            print(f'  row {row}: {d} @ {p_:.2f} logged {was} -> {now}')
        print()
    return main, disc, orders, px


def occupancy(main, name, sleeve, day):
    """Was this sleeve holding a position on `day`? Open-ended if never sold.
    A position bought and sold the same day did not freeze the next morning."""
    for t in main:
        if t['name'] != name or t['sleeve'] != sleeve or not t['buy']:
            continue
        if t['buy'] <= day and (t['sell'] is None or t['sell'] > day):
            return t
    return None


def sessions_between(ser, a, b):
    return sum(1 for d in ser if a < d <= b) if a and b else None


def main_():
    path = sys.argv[1]
    main, disc, orders, px = read_logs(path)
    asof = max(max(s) for s in px.values() if s)
    print(f'workbook through {asof}: {len(main)} model trades, {len(disc)} discretionary\n')

    # --- the live freeze picture -------------------------------------------
    frozen = [nm for nm in sorted({k[0] for k in orders})
              if all(orders.get((nm, s), {}).get('status') == 'HOLDING'
                     for s in ('Bayes', 'OU'))]
    print('MODEL SLEEVE STATE TODAY')
    for nm in sorted({k[0] for k in orders}):
        cells = [orders.get((nm, s), {}) for s in ('Bayes', 'OU')]
        st = '/'.join(str(c.get('status')) for c in cells)
        cl = cells[0].get('close')
        gaps = []
        for c in cells:
            t = c.get('target')
            gaps.append(f'{(t/cl-1)*100:+.1f}%' if (t and cl) else '   -  ')
        notes = ' '.join(str(c.get('note') or '') for c in cells).strip()
        mark = '  <= BOTH SLEEVES FROZEN' if nm in frozen else ''
        print(f'  {nm:5s} {st:18s} close {cl:>8.2f}  to target {gaps[0]:>7s} / {gaps[1]:>7s}'
              f'  {notes:12s}{mark}')
    print(f'\n{len(frozen)} of {len(orders)//2} names have both sleeves frozen: {frozen}')

    # --- every discretionary trade against the occupancy record ------------
    print('\nDISCRETIONARY LOG vs MODEL SLEEVE OCCUPANCY ON THE ENTRY DATE')
    print(f'  {"stock":6s}{"bought":>12s}{"px":>9s}{"sold":>12s}{"ret":>8s}{"sess":>6s}'
          f'{"P&L":>12s}  sleeves held that day')
    rows = []
    for t in disc:
        ser = px.get(t['name'], {})
        held = [s for s in ('Bayes', 'OU') if occupancy(main, t['name'], s, t['buy'])]
        if t['sell'] and t['spx']:
            ret = t['spx'] / t['px'] - 1
            sess = sessions_between(ser, t['buy'], t['sell'])
            sold, pnl, openpos = str(t['sell']), t['pnl'], False
        else:
            last = max(ser) if ser else None
            ret = (ser[last] / t['px'] - 1) if last else None
            sess = sessions_between(ser, t['buy'], last)
            sold, pnl, openpos = 'OPEN', (ser[last] - t['px']) * (t['shares'] or 0) if last else None, True
        rows.append(dict(t, ret=ret, sess=sess, frozen=len(held) == 2, held=held, open=openpos,
                         mark=pnl))
        print(f"  {t['name']:6s}{str(t['buy']):>12s}{t['px']:>9.2f}{sold:>12s}"
              f"{ret*100:>+7.2f}%{sess if sess is not None else 0:>6d}{pnl:>12,.0f}"
              f"  {'+'.join(held) if held else '(none — name was free)'}")

    def block(tag, sel):
        s = [r for r in rows if sel(r)]
        if not s:
            print(f'  {tag:34s} none')
            return None
        rets = [r['ret'] for r in s if r['ret'] is not None]
        sess = [r['sess'] for r in s if r['sess']]
        pnl = sum(r['mark'] or 0 for r in s)
        st = dict(n=len(s), closed=sum(1 for r in s if not r['open']),
                  avg=statistics.fmean(rets), med=statistics.median(rets),
                  worst=min(rets), med_sess=statistics.median(sess) if sess else None,
                  pnl=pnl, losers=sum(1 for r in rets if r < 0))
        print(f"  {tag:34s} n={st['n']:2d} ({st['closed']} closed)  avg {st['avg']*100:+6.2f}%"
              f"  med {st['med']*100:+6.2f}%  worst {st['worst']*100:+7.2f}%"
              f"  med hold {st['med_sess'] or 0:>2.0f} sess  P&L {st['pnl']:>12,.0f}")
        return st

    print('\nSPLIT — the question the proposal turns on:')
    res = {}
    res['frozen'] = block('re-entry, BOTH sleeves frozen', lambda r: r['frozen'])
    res['partial'] = block('one sleeve held', lambda r: len(r['held']) == 1)
    res['free'] = block('name was wholly free', lambda r: not r['held'])
    res['all'] = block('all discretionary', lambda r: True)

    with open(OUT, 'w') as f:
        json.dump(dict(asof=str(asof), frozen_today=frozen,
                       trades=[{k: (str(v) if isinstance(v, dt.date) else v)
                                for k, v in r.items()} for r in rows],
                       summary=res), f, indent=1, default=str)
    print(f'\nsaved {OUT}')


if __name__ == '__main__':
    main_()
