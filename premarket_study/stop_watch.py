"""
stop_watch.py — which open positions are near the 50-calendar-day stop?

The stop is the ONLY mechanism that turns an unrealised drawdown into a
realised loss: a resting take-profit either fills or does not, but at
entry + 50 CALENDAR days the sleeve sells AT THE OPEN regardless of price
(book_sim STOP_DAYS=50, exit at O). §3.63 measured those exits at 3.9% of all
fills, and found them the only marketable part of the loop.

"Release price" is the take-profit that would close the position profitably
BEFORE the stop fires. The engine sets target = bid + prev_close*premium; with
bid ~ prev_close that is ~ buy*(1+premium), which is what is shown — an
indication, not the sheet's own figure.

Reads the blotter from the uploaded workbook. Usage:
  BAYES_WORKBOOK=<xlsx> python3 stop_watch.py [--asof 2026-10-20]
"""
import datetime as dt
import json
import os
import sys

import openpyxl

STOP_DAYS = 50
DISC_STOP_SESSIONS = 10


def main():
    path = os.environ.get('BAYES_WORKBOOK')
    asof = dt.date.fromisoformat(sys.argv[sys.argv.index('--asof') + 1]) \
        if '--asof' in sys.argv else dt.date.today()
    cfg = json.load(open('deployed_book12.json'))['params']
    wb = openpyxl.load_workbook(path, data_only=True)
    at = wb['Active Trading']

    rows = []
    for r in range(42, 1032):
        nm, buy, sell = at.cell(r, 2).value, at.cell(r, 4).value, at.cell(r, 8).value
        if nm in (None, '') or sell not in (None, ''):
            continue
        d = buy.date() if isinstance(buy, dt.datetime) else buy
        tr = at.cell(r, 3).value
        px, sh = at.cell(r, 5).value, at.cell(r, 6).value
        cost = at.cell(r, 7).value or (px * sh)
        prem = cfg.get(nm, {}).get('premium' if tr == 'Bayes' else 'ou_prem')
        rows.append(dict(row=r, name=nm, tr=tr, buy=d, px=px, sh=sh, cost=cost,
                         held=(asof - d).days, stop=d + dt.timedelta(days=STOP_DAYS),
                         left=STOP_DAYS - (asof - d).days,
                         rel=(px * (1 + prem) if prem else None), prem=prem))
    rows.sort(key=lambda x: x['left'])

    print(f'as of {asof} — blotter from the workbook (last saved ~14-15 Oct, so any '
          f'position\nopened or closed since then is not reflected)\n')
    print(f'  {"name":6s}{"sleeve":7s}{"bought":12s}{"held":>6s}{"stop date":>13s}'
          f'{"days left":>11s}{"cost":>11s}{"buy px":>10s}{"release px":>12s}{"needs":>8s}')
    tot = 0.0
    for x in rows:
        tot += x['cost']
        need = (x['rel'] / x['px'] - 1) * 100 if x['rel'] else float('nan')
        print(f"  {x['name']:6s}{x['tr']:7s}{str(x['buy']):12s}{x['held']:>5d}d"
              f"{str(x['stop']):>13s}{x['left']:>10d}d{x['cost']/1e3:>10.0f}k"
              f"{x['px']:>10.2f}{x['rel']:>12.2f}{need:>7.2f}%")
    print(f'  {"TOTAL":6s}{"":7s}{"":12s}{"":6s}{"":13s}{"":11s}{tot/1e3:>10.0f}k')

    print(f'\n  --- grouped by stop date ---\n')
    by = {}
    for x in rows:
        by.setdefault(x['stop'], []).append(x)
    for sd in sorted(by):
        g = by[sd]
        c = sum(x['cost'] for x in g)
        print(f"  {sd}  in {g[0]['left']:>2d} days   ${c/1e3:>6.0f}k across "
              f"{len(g)} sleeve(s): " + ', '.join(f"{x['name']} {x['tr']}" for x in g))

    print(f'\n  --- discretionary log ({DISC_STOP_SESSIONS}-session time stop) ---\n')
    for r in range(42, 1032):
        nm, buy, sell = at.cell(r, 21).value, at.cell(r, 23).value, at.cell(r, 28).value
        if nm in (None, '') or sell not in (None, ''):
            continue
        d = buy.date() if isinstance(buy, dt.datetime) else buy
        print(f"  r{r} {nm} {at.cell(r,22).value}  bought {d}  "
              f"{(asof-d).days}d ago  {at.cell(r,25).value} sh @ {at.cell(r,24).value}")

    json.dump([{k: str(v) for k, v in x.items()} for x in rows],
              open('stop_watch.json', 'w'), indent=1)
    print('\nwrote stop_watch.json')


if __name__ == '__main__':
    main()
