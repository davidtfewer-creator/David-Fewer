"""
live_record.py — what the live blotter says about each name (7 Oct 2026).

User: "I'm not happy with CF's performance. With optimised parameters CF
relies on a small number of high-margin trades to meet its return. That is the
wrong profile for the model -- it lacks predictability. Fast turnover stocks
are a better fit."

That is a claim about the SHAPE of a name's return, not its level, so it needs
shape statistics: how many trades, how long they are held, and how much of the
name's P&L rides on its best few trades. A name whose top third of trades
carries nearly all its profit is one bad quarter away from nothing; a name
whose P&L is spread evenly compounds predictably.

Reads the live blotter (Active Trading A41:L...) and reports per name:
  trades, realised P&L, average and median return, median hold in sessions,
  the share of P&L from the best third of trades, and the yield per invested
  dollar-day -- the 3.27 measure, which is what the pool actually pays for.

Usage: python live_record.py <workbook.xlsx>
"""
import datetime as dt
import json
import statistics
import sys

import openpyxl

OUT = 'live_record.json'


def d_of(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    at = wb['Active Trading']
    q = wb['Query']
    sessions = sorted({d for d in (d_of(q.cell(row=r, column=1).value)
                                   for r in range(2, q.max_row + 1)) if d})
    trades = []
    for r in range(42, at.max_row + 1):
        nm = at.cell(row=r, column=2).value
        if not nm:
            continue
        buy, sell = d_of(at.cell(row=r, column=4).value), d_of(at.cell(row=r, column=8).value)
        cost, pnl = at.cell(row=r, column=7).value, at.cell(row=r, column=11).value
        if not (buy and sell and isinstance(cost, (int, float))
                and isinstance(pnl, (int, float)) and cost > 0):
            continue
        held = sum(1 for d in sessions if buy < d <= sell)
        trades.append(dict(name=str(nm).strip(), sleeve=str(at.cell(row=r, column=3).value or ''),
                           buy=buy, sell=sell, cost=cost, pnl=pnl, ret=pnl / cost,
                           hold=max(held, 0)))
    wb.close()
    return trades, sessions


def top_third_share(tr):
    """Share of total P&L contributed by the best third of trades. 0.33 means
    perfectly even; near 1.0 means the name is a few trades in a trenchcoat."""
    pos = sorted((t['pnl'] for t in tr), reverse=True)
    tot = sum(pos)
    if tot <= 0 or len(pos) < 3:
        return None
    k = max(1, round(len(pos) / 3))
    return sum(pos[:k]) / tot


def stats(tr):
    rets = [t['ret'] for t in tr]
    dd = [max(t['hold'], 1) for t in tr]
    return dict(n=len(tr), pnl=sum(t['pnl'] for t in tr),
                avg=statistics.fmean(rets), med=statistics.median(rets),
                worst=min(rets), losers=sum(1 for r in rets if r < 0),
                med_hold=statistics.median(dd),
                dollar_day=sum(t['pnl'] for t in tr) / sum(t['cost'] * max(t['hold'], 1)
                                                           for t in tr),
                top3=top_third_share(tr))


def main():
    trades, sessions = load(sys.argv[1])
    span = (max(t['sell'] for t in trades) - min(t['buy'] for t in trades)).days
    print(f'{len(trades)} closed live trades, {min(t["buy"] for t in trades)} to '
          f'{max(t["sell"] for t in trades)} ({span} days)\n')
    print(f'  {"name":6s}{"n":>4s}{"P&L":>12s}{"avg":>8s}{"med":>8s}{"worst":>8s}'
          f'{"lose":>6s}{"hold":>6s}{"%/$-day":>9s}{"top3rd":>8s}')
    byname = {}
    for t in trades:
        byname.setdefault(t['name'], []).append(t)
    res = {}
    for nm in sorted(byname, key=lambda n: -len(byname[n])):
        s = stats(byname[nm])
        res[nm] = s
        print(f'  {nm:6s}{s["n"]:>4d}{s["pnl"]:>12,.0f}{s["avg"]*100:>+7.2f}%'
              f'{s["med"]*100:>+7.2f}%{s["worst"]*100:>+7.2f}%{s["losers"]:>6d}'
              f'{s["med_hold"]:>6.0f}{s["dollar_day"]*100:>8.3f}%'
              f'{(s["top3"]*100 if s["top3"] else 0):>7.0f}%')
    s = stats(trades)
    res['BOOK'] = s
    print(f'  {"BOOK":6s}{s["n"]:>4d}{s["pnl"]:>12,.0f}{s["avg"]*100:>+7.2f}%'
          f'{s["med"]*100:>+7.2f}%{s["worst"]*100:>+7.2f}%{s["losers"]:>6d}'
          f'{s["med_hold"]:>6.0f}{s["dollar_day"]*100:>8.3f}%'
          f'{(s["top3"]*100 if s["top3"] else 0):>7.0f}%')

    print('\n  turnover: trades per name per 100 sessions of the live span')
    n_sess = sum(1 for d in sessions
                 if min(t['buy'] for t in trades) <= d <= max(t['sell'] for t in trades))
    for nm in sorted(byname, key=lambda n: -len(byname[n])):
        print(f'    {nm:6s}{len(byname[nm])/n_sess*100:>6.1f} trades/100 sessions'
              f'   (sleeve-days held: {sum(max(t["hold"],1) for t in byname[nm])})')

    with open(OUT, 'w') as f:
        json.dump(res, f, indent=1, default=str)
    print(f'\nsaved {OUT}')


if __name__ == '__main__':
    main()
